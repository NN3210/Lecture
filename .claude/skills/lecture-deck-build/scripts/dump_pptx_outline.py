import sys
sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")
from pptx import Presentation
from pptx.util import Emu

if len(sys.argv) < 2 or sys.argv[1].startswith("-"):
    sys.exit("使い方: python -I dump_pptx_outline.py <pptx>   全枚の見出し・本文・表を Markdown に出す（> で保存）")
path = sys.argv[1]
prs = Presentation(path)
print(f"# {path}")
print(f"slide size: {prs.slide_width} x {prs.slide_height} EMU")
print(f"slides: {len(prs.slides)}")
for i, slide in enumerate(prs.slides, 1):
    hidden = slide._element.get("show") == "0"
    layout = slide.slide_layout.name
    title = ""
    if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
        title = slide.shapes.title.text_frame.text.strip().replace("\n", " / ")
    print(f"\n## [{i:03d}] {'(HIDDEN) ' if hidden else ''}layout={layout} title={title!r}")
    npic = 0
    ntbl = 0
    for sh in slide.shapes:
        kind = sh.shape_type
        if sh.has_text_frame:
            txt = sh.text_frame.text.strip()
            if txt and (slide.shapes.title is None or sh.shape_id != slide.shapes.title.shape_id):
                txt1 = txt.replace("\n", " | ")
                if len(txt1) > 400:
                    txt1 = txt1[:400] + "…"
                print(f"  - text: {txt1}")
        if sh.shape_type == 13:
            npic += 1
        if getattr(sh, "has_table", False) and sh.has_table:
            ntbl += 1
            rows = []
            for r in sh.table.rows:
                rows.append(" | ".join(c.text.strip().replace("\n", " ") for c in r.cells))
            print("  - table: " + " || ".join(rows)[:400])
        if sh.shape_type == 6:  # group
            for g in sh.shapes:
                if g.has_text_frame and g.text_frame.text.strip():
                    print(f"  - grp-text: {g.text_frame.text.strip().replace(chr(10),' | ')[:200]}")
                if g.shape_type == 13:
                    npic += 1
    # count math objects
    xml = slide._element.xml
    nmath = xml.count("<m:oMath>") + xml.count("<m:oMath ")
    print(f"  - meta: pictures={npic} tables={ntbl} math={nmath}")
    if slide.has_notes_slide:
        nt = slide.notes_slide.notes_text_frame.text.strip()
        if nt:
            print(f"  - notes: {nt.replace(chr(10),' | ')[:300]}")
