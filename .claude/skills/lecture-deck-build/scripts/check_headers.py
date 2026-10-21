#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""pptx の全枚からヘッダー文字列とページ番号を抜き出し、非表示・番号の不一致・章番号の戻りを数える。
PNG を全枚見る前の機械点検。

使い方:
    python -I $SK/lecture-deck-build/scripts/check_headers.py <pptx> --header-prefix "測量学 第3回" [--quiet]
"""
from __future__ import annotations

import argparse
import re
import sys

from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8")

PAGE = re.compile(r"^\s*(\d{1,3})\s*/\s*(\d{2,3})\s*$")
BARE = re.compile(r"^\s*(\d{1,3})\s*$")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx")
    ap.add_argument("--header-prefix", required=True)
    ap.add_argument("--quiet", action="store_true", help="枚ごとの行を出さず、集計だけ出す")
    a = ap.parse_args()
    prs = Presentation(a.pptx)
    total = len(prs.slides)
    foot = int(prs.slide_height * 0.85)
    hidden, mismatch, no_header, no_number, back = [], [], [], [], []
    prev_ch = None
    for i, sl in enumerate(prs.slides, 1):
        is_hidden = sl._element.get("show") == "0"
        header, page = "", ""
        for sh in sl.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                t = "".join(r.text for r in p.runs).strip()
                if t.startswith(a.header_prefix) and not header:
                    header = t
                elif PAGE.match(t):
                    page = t
                elif BARE.match(t) and sh.top is not None and sh.top >= foot and not page:
                    page = t
        if is_hidden:
            hidden.append(i)
        if not header:
            no_header.append(i)
        if not page:
            no_number.append(i)
        else:
            m = PAGE.match(page) or BARE.match(page)
            if int(m.group(1)) != i:
                mismatch.append(f"{i}({page})")
        ch = None
        if header:
            m = re.search(r"/\s*(\d{2})\s*\|", header)
            if m:
                ch = int(m.group(1))
                if prev_ch is not None and ch < prev_ch:
                    back.append(f"{i}({prev_ch:02d}→{ch:02d})")
                prev_ch = ch
        if not a.quiet:
            print(f"{i:03d} | {header or '—'} | {page or '—'}{' | (HIDDEN)' if is_hidden else ''}")
    print(f"\n枚数 {total}。非表示 {len(hidden)} {hidden or ''}")
    print(f"ページ番号が位置と違う {len(mismatch)} {mismatch or ''}")
    print(f"章番号が戻る {len(back)} {back or ''}")
    print(f"ヘッダー無し {len(no_header)}（旧世代の枚など）{no_header if len(no_header) <= 20 else str(no_header[:20]) + '…'}")
    print(f"番号無し {len(no_number)}{no_number if len(no_number) <= 20 else str(no_number[:20]) + '…'}")


if __name__ == "__main__":
    main()
