#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""既存デッキの枚を「構成案」の順に並べ替え、使わない枚を削除し、非表示を解除し、
橋渡しの新規枚を既存枚を型紙として複製して差し込む。PowerPoint は使わない（python-pptx のみ）。

使い方（科目のプロジェクト直下で実行）:
    python -I $SK/lecture-deck-build/scripts/build_pptx.py <元.pptx> <構成案.md> <出力.pptx>
        [--new-spec <新規枚.md>] [--renumber] [--header-prefix "測量学 第3回"] [--shape-map <map.json>]
    python -I .../build_pptx.py <元.pptx> --list-shapes 25      型紙候補の枚の部品名と文字を出す

構成案（Markdown 表）の読み方:
    行の 2 列目が「k1:NN」（接頭辞は英小文字なら何でもよい）なら元デッキの NN 枚目（1 始まり、非表示を含む並び順）、
    「new:ID」なら新規枚。
    | 順 | 元 | 題 | 備考 |
    | 1 | k1:1 | 表紙 | |
    | 2 | new:goal | 今日のゴール | 新規 |
    | 3 | k1:22 | みかんの皮 | 章を移した。header=測量学 第3回  /  03  \\|  地球を地図にする |
    備考欄の「header=…」は、その枚の「<ヘッダー接頭辞> /」で始まる段落を書き換える（表の中なので | は \\| と書く）。
    --header-prefix を省くと、最初の header= の値の「/」より前を接頭辞にする。

新規枚の仕様（Markdown）:
    ## new:goal
    - template: 25            ← 型紙にする元デッキの枚（既定の部品名は 3 枚カード型）
    - header: 測量学 第3回  /  00  |  今日のゴール
    - title: 今日のゴール
    - subtitle: 1 行の副題（省略可）
    - card1: 見出し | 本文 1 行目 / 2 行目      ← card1〜card3
    - bullets: 1 つ目 / 2 つ目 / 3 つ目           ← 箇条書き型（カードを消して 1 つの枠にする）
    - takeaway: 下の帯の 1 文
    - notes: 発表者ノート（省略可）
"""
from __future__ import annotations

import argparse
import copy
import json
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

sys.stdout.reconfigure(encoding="utf-8")

NS_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
ROW = re.compile(r"^\|\s*\d+\s*\|\s*(?:new:([A-Za-z0-9_\-]+)|([a-z][a-z0-9]*):(\d+))\s*\|")

# 既定の型紙（第 3 回・視覚版の 3 枚カード型）の部品名
DEFAULT_MAP = {
    "header": "TextBox 2", "title": "TextBox 3", "subtitle": "TextBox 4", "page": "TextBox 5",
    "cards": [["TextBox 9", "TextBox 10"], ["TextBox 14", "TextBox 15"], ["TextBox 19", "TextBox 20"]],
    "card_shapes": ["Rounded Rectangle 6", "Rounded Rectangle 7", "TextBox 8", "TextBox 9", "TextBox 10",
                    "Rounded Rectangle 11", "Rounded Rectangle 12", "TextBox 13", "TextBox 14", "TextBox 15",
                    "Rounded Rectangle 16", "Rounded Rectangle 17", "TextBox 18", "TextBox 19", "TextBox 20"],
    "takeaway": "TextBox 23",
}


def read_plan(path: Path) -> tuple[list[tuple[str, str]], dict[int, str]]:
    """構成案の行を読む。("src", "12") か ("new", "id") の列と、位置→ヘッダー文字列の辞書を返す。"""
    plan: list[tuple[str, str]] = []
    headers: dict[int, str] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        m = ROW.match(line)
        if not m:
            continue
        plan.append(("new", m.group(1)) if m.group(1) else ("src", m.group(3)))
        hm = re.search(r"header=((?:\\\||[^|])+)", line[m.end():])
        if hm:
            headers[len(plan) - 1] = hm.group(1).replace("\\|", "|").strip()
    return plan, headers


def replace_para_text(slide, pat: re.Pattern, new_text: str, min_top: int | None = None) -> int:
    """pat に合う段落すべてで、1 つ目の run に new_text を入れ、他の run を空にする（書式は 1 つ目の run のまま）。
    min_top を与えると、上端がそれより下にある枠だけを対象にする（脚注の番号だけを狙うため）。"""
    n = 0
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if min_top is not None and (shape.top is None or shape.top < min_top):
            continue
        for para in shape.text_frame.paragraphs:
            txt = "".join(r.text for r in para.runs)
            if pat.match(txt) and para.runs:
                para.runs[0].text = new_text
                for r in para.runs[1:]:
                    r.text = ""
                n += 1
    return n


def read_spec(path: Path) -> dict[str, dict]:
    spec: dict[str, dict] = {}
    cur = None
    for line in path.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^##\s+new:([A-Za-z0-9_\-]+)", line)
        if m:
            cur = spec.setdefault(m.group(1), {})
            continue
        m = re.match(r"^-\s+([a-z0-9]+):\s*(.*)$", line)
        if m and cur is not None:
            cur[m.group(1)] = m.group(2).strip()
    return spec


def shape_by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(f"型紙に部品 {name!r} が無い（--list-shapes で部品名を確かめ、--shape-map で渡す）")


def set_text(shape, lines: list[str]) -> None:
    """最初の段落・最初のランの書式を残して、文を入れ替える。"""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    rpr = copy.deepcopy(p0.runs[0]._r.rPr) if p0.runs and p0.runs[0]._r.rPr is not None else None
    ppr = copy.deepcopy(p0._p.pPr) if p0._p.pPr is not None else None
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    for i, text in enumerate(lines):
        p = p0 if i == 0 else tf.add_paragraph()
        if i > 0 and ppr is not None:
            p._p.insert(0, copy.deepcopy(ppr))
        r = p.add_run()
        r.text = text
        if rpr is not None:
            r._r.insert(0, copy.deepcopy(rpr))


def clone_slide(prs, src_slide):
    """型紙の枚を複製して末尾に追加する（図は無い前提。関係のコピーはしない）。"""
    layout = src_slide.slide_layout
    new = prs.slides.add_slide(layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    for el in src_slide.shapes._spTree:
        if el.tag.endswith("}sp") or el.tag.endswith("}cxnSp") or el.tag.endswith("}grpSp"):
            new.shapes._spTree.append(copy.deepcopy(el))
    return new


def build_new_slide(prs, spec: dict, sid: str, names: dict):
    tnum = int(spec.get("template", "25"))
    tmpl = prs.slides[tnum - 1]
    s = clone_slide(prs, tmpl)
    if "header" in spec:
        set_text(shape_by_name(s, names["header"]), [spec["header"]])
    if "title" in spec:
        set_text(shape_by_name(s, names["title"]), [spec["title"]])
    if "subtitle" in spec:
        set_text(shape_by_name(s, names["subtitle"]), [spec["subtitle"]])
    elif names.get("subtitle"):
        set_text(shape_by_name(s, names["subtitle"]), [""])
    if "bullets" in spec:
        # カードを消して、カード 1 の本文枠を広げて箇条書きにする
        body_name = names["cards"][0][1]
        body = shape_by_name(s, body_name)
        for nm in names["card_shapes"]:
            if nm != body_name:
                el = shape_by_name(s, nm)._element
                el.getparent().remove(el)
        body.left, body.top, body.width, body.height = Inches(0.9), Inches(2.3), Inches(11.6), Inches(3.9)
        set_text(body, ["・" + b.strip() for b in spec["bullets"].split(" / ")])
    else:
        cards = [spec.get(f"card{i}") for i in (1, 2, 3)]
        for (hname, bname), c in zip(names["cards"], cards):
            if not c:
                continue
            head, _, body = c.partition("|")
            set_text(shape_by_name(s, hname), [head.strip()])
            set_text(shape_by_name(s, bname), [ln.strip() for ln in body.split(" / ") if ln.strip()])
    if "takeaway" in spec:
        set_text(shape_by_name(s, names["takeaway"]), [spec["takeaway"]])
    if "notes" in spec:
        s.notes_slide.notes_text_frame.text = spec["notes"]
    s.name = f"new:{sid}"
    return s


def list_shapes(prs, n: int) -> None:
    sl = prs.slides[n - 1]
    print(f"[{n:03d}] layout={sl.slide_layout.name} hidden={sl._element.get('show') == '0'}")
    for sh in sl.shapes:
        txt = sh.text_frame.text.strip().replace("\n", " | ")[:60] if sh.has_text_frame else ""
        print(f"  {sh.name!r:28} type={sh.shape_type} top={sh.top} left={sh.left} text={txt!r}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("src")
    ap.add_argument("plan", nargs="?")
    ap.add_argument("dst", nargs="?")
    ap.add_argument("--new-spec", default="")
    ap.add_argument("--renumber", action="store_true",
                    help="ページ番号を位置に書き換える（「NN / 81」→「位置 / 総数」、下端 15%% の数字だけの脚注「14」→「位置」）")
    ap.add_argument("--header-prefix", default="", help="ヘッダー段落の接頭辞（例「測量学 第3回」）。省くと header= の値から取る")
    ap.add_argument("--shape-map", default="", help="型紙の部品名を書いた JSON（既定は 3 枚カード型）")
    ap.add_argument("--list-shapes", type=int, default=0, metavar="N", help="元デッキの N 枚目の部品名を出して終わる")
    a = ap.parse_args()

    prs = Presentation(a.src)
    if a.list_shapes:
        list_shapes(prs, a.list_shapes)
        return
    if not a.plan or not a.dst:
        sys.exit("構成案と出力先を指定する（または --list-shapes N）")

    plan, headers = read_plan(Path(a.plan))
    if not plan:
        sys.exit("構成案に <接頭辞>:NN / new:ID の行が無い")
    spec = read_spec(Path(a.new_spec)) if a.new_spec else {}
    names = json.loads(Path(a.shape_map).read_text(encoding="utf-8")) if a.shape_map else DEFAULT_MAP
    prefix = a.header_prefix.strip()
    if not prefix and headers:
        prefix = next(iter(headers.values())).split("/")[0].strip()
        print(f"ヘッダー接頭辞を header= から推定: {prefix!r}")
    header_pat = re.compile(r"^\s*" + re.escape(prefix) + r"\s*/") if prefix else None

    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    n_src = len(ids)

    used = [int(v) for k, v in plan if k == "src"]
    bad = [n for n in used if n < 1 or n > n_src]
    if bad:
        sys.exit(f"元デッキに無い枚番号: {bad}（元は {n_src} 枚）")
    dup = sorted({n for n in used if used.count(n) > 1})
    if dup:
        sys.exit(f"同じ枚が 2 回以上使われている: {dup}")
    missing_spec = [v for k, v in plan if k == "new" and v not in spec]
    if missing_spec:
        sys.exit(f"新規枚の仕様が無い: {missing_spec}")

    new_ids: dict[str, object] = {}
    for k, v in plan:
        if k == "new":
            build_new_slide(prs, spec[v], v, names)
            new_ids[v] = list(sldIdLst)[-1]

    order = [ids[int(v) - 1] if k == "src" else new_ids[v] for k, v in plan]
    for el in list(sldIdLst):
        sldIdLst.remove(el)
    for el in order:
        sldIdLst.append(el)

    keep = set(used)
    removed = 0
    for n, el in enumerate(ids, 1):
        if n not in keep:
            prs.part.drop_rel(el.get(NS_R))
            removed += 1

    for slide in prs.slides:
        if slide._element.get("show") == "0":
            del slide._element.attrib["show"]

    n_hdr = 0
    if headers and header_pat is None:
        print("警告: header= があるがヘッダー接頭辞が決まらない（--header-prefix を指定）")
    for pos, slide in enumerate(prs.slides):
        if pos in headers and header_pat is not None:
            if replace_para_text(slide, header_pat, headers[pos]):
                n_hdr += 1
            else:
                print(f"警告: {pos + 1} 枚目にヘッダー文字列が見つからない（header= は無視）")

    if a.renumber:
        total = len(plan)
        pat = re.compile(r"^\s*\d{1,3}\s*/\s*\d{2,3}\s*$")
        bare = re.compile(r"^\s*\d{1,3}\s*$")
        foot = int(prs.slide_height * 0.85)
        n_num = 0
        for pos, slide in enumerate(prs.slides, 1):
            n_num += replace_para_text(slide, pat, f"{pos:02d} / {total}")
            n_num += replace_para_text(slide, bare, str(pos), min_top=foot)
        print(f"ページ番号書換 {n_num} 段落（番号の無い枚はそのまま）")

    prs.save(a.dst)
    print(f"元 {n_src} 枚 → 使用 {len(used)} 枚、削除 {removed} 枚、新規 {len(new_ids)} 枚、"
          f"ヘッダー書換 {n_hdr} 枚、出力 {len(prs.slides)} 枚: {a.dst}")


if __name__ == "__main__":
    main()
