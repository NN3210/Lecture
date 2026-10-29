#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Marp 原稿から2026年AI講義用の差込PPTXを再生成する。

使用例: python scripts/build_ai_2026.py
唯一の本文・出典・動画指定は各 materials/slides/.../slides.md に置く。
"""
from __future__ import annotations

import io
import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SETS = (
    (ROOT / "materials/slides/2026_AI農業更新", "AI農業更新_差込用_2026.pptx"),
    (ROOT / "materials/slides/2026_AI解説補足", "AI解説補足_2026.pptx"),
)
W, H = 13.333333, 7.5
INK = RGBColor(37, 49, 41)
MUTED = RGBColor(82, 95, 85)
GREEN = RGBColor(69, 112, 76)
PALE = RGBColor(233, 241, 232)
PAPER = RGBColor(252, 253, 250)
LINE = RGBColor(185, 203, 184)
FONT = "Yu Gothic"


def parse(path: Path):
    text = path.read_text(encoding="utf-8")
    text = re.sub(r"\A---\n.*?\n---\n", "", text, flags=re.S)
    slides = []
    for raw in re.split(r"^---\s*$", text, flags=re.M):
        if not raw.strip():
            continue
        directives = dict(re.findall(r"<!--\s*_(\w+):\s*(.*?)\s*-->", raw))
        title = re.search(r"^#\s+(.+)$", raw, flags=re.M)
        if not title:
            raise ValueError(f"スライドに見出しがありません: {raw[:80]}")
        sections = []
        clean = re.sub(r"<!--.*?-->", "", raw, flags=re.S)
        for block in re.split(r"(?=^##\s+)", clean, flags=re.M):
            lines = [x.strip() for x in block.strip().splitlines() if x.strip()]
            if lines and lines[0].startswith("## "):
                sections.append((lines[0][3:], " ".join(lines[1:])))
        slides.append({"title": title.group(1).strip(), "sections": sections, **directives})
    return slides


def shape_text(slide, text, x, y, w, h, size, color=INK, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.02)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return tb


def box(slide, x, y, w, h, fill=PALE, line=None, radius=False):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    effect = s._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef")
    if effect is not None:
        effect.set("idx", "0")
    return s


def base_slide(prs, data, num, deck_tag):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PAPER
    box(slide, 0, 0, .17, H, GREEN)
    shape_text(slide, deck_tag, .56, .13, 6.8, .3, 10, GREEN, True)
    shape_text(slide, data["title"], .56, .52, 12.1, .9, 28, INK, True)
    box(slide, .56, 1.46, 12.1, .025, LINE)
    shape_text(slide, f"2026.10.30  |  {num:02d}", 10.73, 7.14, 1.94, .2, 8, MUTED, align=PP_ALIGN.RIGHT)
    notes = data.get("notes", "")
    sources = data.get("sources", "")
    if notes or sources:
        slide.notes_slide.notes_text_frame.text = f"話す内容: {notes}\n\n出典: {sources}"
    return slide


def add_section(slide, label, body, x, y, w, h, accent=False, index=None):
    fill = GREEN if accent else PALE
    box(slide, x, y, w, h, fill)
    c1 = RGBColor(255, 255, 255) if accent else GREEN
    c2 = RGBColor(255, 255, 255) if accent else INK
    if index is not None:
        shape_text(slide, f"{index:02d}", x+.18, y+.12, .48, .36, 13, c1, True)
        lx = x+.72
    else:
        lx = x+.24
    shape_text(slide, label, lx, y+.15, w-(lx-x)-.2, .35, 16, c1, True)
    shape_text(slide, body, x+.24, y+.61, w-.48, h-.75, 20 if len(body)<35 else 18, c2)


def draw_flow(slide, sections):
    n = len(sections)
    if n == 4:
        xs = [.65, 3.77, 6.89, 10.01]
        width = 2.69
        for i, ((a,b), x) in enumerate(zip(sections, xs)):
            add_section(slide, a, b, x, 2.18, width, 3.8, accent=(i==n-1), index=i+1)
            if i<3:
                shape_text(slide, "→", x+width+.13, 3.61, .27, .4, 20, GREEN, True, PP_ALIGN.CENTER)
    else:
        ys=[1.85,3.15,4.45,5.75]
        for i,(a,b) in enumerate(sections[:4]):
            add_section(slide,a,b,.82,ys[i],11.7,1.06,accent=(i==3),index=i+1)


def draw_case(slide, sections):
    n=len(sections)
    for i,(a,b) in enumerate(sections):
        y=1.86+i*1.18
        accent=i==n-1
        box(slide,.78,y,11.86,1.0,GREEN if accent else PALE)
        c=RGBColor(255,255,255) if accent else GREEN
        shape_text(slide,a,1.04,y+.18,2.2,.65,19,c,True)
        shape_text(slide,b,3.38,y+.17,8.75,.68,21,RGBColor(255,255,255) if accent else INK)


def draw_compare(slide, sections):
    if len(sections)<2:
        return
    a,b=sections[:2]
    add_section(slide,a[0],a[1],.84,1.91,5.7,3.28)
    add_section(slide,b[0],b[1],6.78,1.91,5.7,3.28,accent=True)
    if len(sections)>2:
        label,body=sections[2]
        shape_text(slide,label,1.05,5.55,1.8,.4,15,GREEN,True)
        shape_text(slide,body,2.7,5.42,9.7,.72,17,INK)


def draw_hierarchy(slide, sections):
    layers=[(.82,1.76,11.7,4.85,PALE),
            (1.89,2.59,9.57,3.2,RGBColor(205,228,204)),
            (3.11,3.5,7.13,1.62,GREEN)]
    for i,((a,b),(x,y,w,h,fill)) in enumerate(zip(sections[:3],layers)):
        box(slide,x,y,w,h,fill)
        c=RGBColor(255,255,255) if i==2 else INK
        shape_text(slide,a,x+.22,y+.12,2.0,.39,21,c,True)
        if i<2:
            shape_text(slide,b,x+2.2,y+.13,w-2.5,.4,16,c)
        else:
            shape_text(slide,b,x+.25,y+.65,w-.5,.68,17,c)


def draw_sprayer(slide, sections):
    # 原稿の4段階をネイティブ図形に変換。外部の機械写真は使用しない。
    xs=[.83,4.51,8.19]
    for i,((a,b),x) in enumerate(zip(sections[:3],xs)):
        add_section(slide,a,b,x,1.82,3.18,1.37,accent=(i==2),index=i+1)
        if i<2:
            shape_text(slide,"→",x+3.25,2.28,.32,.36,22,GREEN,True,PP_ALIGN.CENTER)
    # カメラ、処理器、ノズルの小さな模式図
    box(slide,1.98,3.48,.72,.4,INK)
    lens=slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL,Inches(2.23),Inches(3.55),Inches(.26),Inches(.26))
    lens.fill.solid();lens.fill.fore_color.rgb=PALE;lens.line.fill.background()
    box(slide,5.71,3.44,.7,.48,PALE,GREEN)
    shape_text(slide,"AI",5.77,3.52,.58,.28,13,GREEN,True,PP_ALIGN.CENTER)
    box(slide,9.42,3.46,.74,.18,INK)
    box(slide,9.69,3.64,.2,.29,GREEN)
    shape_text(slide,"↓",9.52,3.89,.52,.44,24,GREEN,True,PP_ALIGN.CENTER)
    # 畝：作物と雑草を区別する構成要素
    box(slide,1.11,5.39,11.1,.035,LINE)
    for x in (2.05,4.1,6.15,8.2,10.25):
        box(slide,x,5.03,.075,.35,GREEN)
        for off in (-.21,.08):
            leaf=slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL,Inches(x+off),Inches(4.89),Inches(.28),Inches(.17))
            leaf.fill.solid();leaf.fill.fore_color.rgb=GREEN;leaf.line.fill.background()
    # 対象の雑草だけを明示
    wx=9.74
    box(slide,wx,5.04,.055,.34,RGBColor(171,91,57))
    for off in (-.16,.05):
        leaf=slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL,Inches(wx+off),Inches(4.98),Inches(.2),Inches(.12))
        leaf.fill.solid();leaf.fill.fore_color.rgb=RGBColor(171,91,57);leaf.line.fill.background()
    shape_text(slide,"作物",1.91,5.56,1.1,.33,16,GREEN,True)
    shape_text(slide,"雑草：ここへ散布",9.02,5.56,2.62,.35,16,RGBColor(150,70,45),True)
    if len(sections)>3:
        shape_text(slide,f"{sections[3][0]}｜{sections[3][1]}",1.12,6.22,11.14,.48,16,MUTED)


def draw_timeline(slide, sections):
    box(slide,.99,3.45,11.35,.04,LINE)
    xs=[.74,3.78,6.82,9.86]
    for i,((a,b),x) in enumerate(zip(sections,xs)):
        box(slide,x+1.16,3.31,.27,.27,GREEN)
        shape_text(slide,a,x,2.22,2.58,.73,26,GREEN,True,PP_ALIGN.CENTER)
        shape_text(slide,b,x,3.83,2.58,1.82,18,INK,align=PP_ALIGN.CENTER)


def make_poster(title):
    im=Image.new("RGB",(1200,600),(35,55,42))
    d=ImageDraw.Draw(im)
    d.rounded_rectangle((40,40,1160,560),radius=24,outline=(120,160,126),width=4)
    d.ellipse((504,204,696,396),fill=(231,241,229))
    d.polygon([(580,250),(580,350),(660,300)],fill=(55,98,62))
    font_path=Path("C:/Windows/Fonts/YuGothB.ttc")
    if not font_path.exists():
        font_path=Path("C:/Windows/Fonts/meiryo.ttc")
    font=ImageFont.truetype(str(font_path),33) if font_path.exists() else ImageFont.load_default()
    d.text((600,480),title,fill=(235,244,233),font=font,anchor="mm")
    stream=io.BytesIO();im.save(stream,format="PNG");stream.seek(0)
    return stream


def add_online_video(slide, url, title):
    m=re.search(r"(?:v=|youtu\.be/)([A-Za-z0-9_-]{11})",url)
    if not m: raise ValueError(f"YouTube URLではありません: {url}")
    embed=f"https://www.youtube.com/embed/{m.group(1)}"
    poster=make_poster(title)
    pic=slide.shapes.add_picture(poster,Inches(1.06),Inches(1.72),width=Inches(11.22),height=Inches(3.78))
    pic.name="公式オンライン動画"
    pic._element.nvPicPr.cNvPr.set("title",title)
    nv=pic._element.nvPicPr.nvPr
    video=OxmlElement("a:videoFile")
    rid=slide.part.relate_to(embed,RT.VIDEO,is_external=True)
    video.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}link",rid)
    nv.append(video)
    hlink=OxmlElement("a:hlinkClick")
    hlink.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id","")
    hlink.set("action","ppaction://media")
    pic._element.nvPicPr.cNvPr.append(hlink)
    return pic


def draw_video(slide, sections, url):
    title=sections[0][1] if sections else "公式動画"
    add_online_video(slide,url,title)
    if sections:
        shape_text(slide,f"{sections[0][0]}  {sections[0][1]}",1.08,5.64,11.2,.36,17,INK,True)
    if len(sections)>1:
        shape_text(slide,f"{sections[1][0]}：{sections[1][1]}",1.08,6.02,11.2,.37,15,MUTED)
    shape_text(slide,"オンライン動画（接続が必要）",1.1,6.57,5.4,.34,13,GREEN,True)
    link=shape_text(slide,"公式動画をブラウザーで開く ↗",7.0,6.53,5.2,.43,14,GREEN,True,PP_ALIGN.RIGHT)
    link.click_action.hyperlink.address=url


def build(folder: Path, filename: str):
    slides=parse(folder/"slides.md")
    prs=Presentation()
    prs.slide_width,prs.slide_height=Inches(W),Inches(H)
    tag="農業AI｜2026年更新" if "農業" in str(folder) else "AIの基礎｜2026年補足"
    for idx,data in enumerate(slides,1):
        sl=base_slide(prs,data,idx,tag)
        kind=data.get("class","flow")
        if kind=="flow":draw_flow(sl,data["sections"])
        elif kind=="case":draw_case(sl,data["sections"])
        elif kind=="compare":draw_compare(sl,data["sections"])
        elif kind=="hierarchy":draw_hierarchy(sl,data["sections"])
        elif kind=="timeline":draw_timeline(sl,data["sections"])
        elif kind=="sprayer":draw_sprayer(sl,data["sections"])
        elif kind=="video":draw_video(sl,data["sections"],data["video"])
        else:raise ValueError(f"未対応レイアウト: {kind}")
    dest=folder/filename
    prs.save(dest)
    print(f"{dest}: {len(slides)}枚")


if __name__=="__main__":
    for f,n in SETS:build(f,n)
