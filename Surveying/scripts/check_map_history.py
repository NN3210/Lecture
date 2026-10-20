#!/usr/bin/env python3
"""完成PDFの全ページ描画・可視文字の照合・保存形式の検証。"""
from pathlib import Path
import json
import re
import hashlib
import zipfile
import fitz
from pptx import Presentation
from PIL import Image, ImageDraw

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"materials/slides/第03回_測量の座標/補足_地図と図法の歴史_20261020"
QA=ROOT/"notes/map_history_20261020"
PPT=OUT/"地図と図法の歴史_補足スライド.pptx"
PDF=PPT.with_suffix(".pdf")
with zipfile.ZipFile(PPT) as z:
    assert z.testzip() is None
prs=Presentation(PPT)
doc=fitz.open(PDF)
assert len(doc)==len(prs.slides)==20
canvas=Image.new("RGB",(1600,1250),"#ddd")
outside=[]
missing=[]
fold=lambda t:re.sub(r"\s+","",t)
for i,p in enumerate(doc):
    pix=p.get_pixmap(matrix=fitz.Matrix(1.2,1.2))
    im=Image.frombytes("RGB",[pix.width,pix.height],pix.samples)
    im.save(QA/f"slide-{i+1:02}.png")
    im.thumbnail((390,220))
    x=(i%4)*400+5;y=(i//4)*250+20
    canvas.paste(im,(x,y))
    ImageDraw.Draw(canvas).text((x,y-16),str(i+1),fill="black")
    for b in p.get_text("dict")["blocks"]:
        for ln in b.get("lines",[]):
            for sp in ln["spans"]:
                x0,y0,x1,y1=sp["bbox"]
                if x0<0 or y0<0 or x1>p.rect.width or y1>p.rect.height:
                    outside.append([i+1,sp["text"]])
    pdftext=fold(p.get_text())
    for a in prs.slides[i].shapes:
        if not a.has_text_frame:continue
        for par in a.text_frame.paragraphs:
            if par.text and fold(par.text) not in pdftext:
                missing.append([i+1,par.text])
    assert prs.slides[i].has_notes_slide
canvas.save(QA/"全20枚プレビュー.jpg",quality=90)
result={"pages":len(doc),"outside_page_text":outside,"pptx_text_missing_in_pdf":missing,
        "zip_integrity":True,"notes_on_every_slide":True,
        "page_size":list(doc[0].rect),
        "files":{p.name:{"bytes":p.stat().st_size,"sha256":hashlib.sha256(p.read_bytes()).hexdigest()} for p in [PPT,PDF]}}
(QA/"PDF検証.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps(result,ensure_ascii=False,indent=2))
assert not outside and not missing
