#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""発表者ノートからナレーション台本（narration.md）の下書きを作る。

入力は次のどちらでもよい。
  - pptx（発表者ノートを読む）
  - Marp 形式の Markdown（slides.md。各スライドの <!-- --> を発表者ノートとして読む。
    枚数の数え方は scripts/md2pptx.py と同じなので、md2pptx.py → pptx_export.ps1 で作った PNG の番号と一致する）

使い方:
  python scripts/ondemand/extract_notes.py <slides.pptx | slides.md> -o narration.md

出力形式（build_video.py が読む）:
  # <タイトル>
  ## 1
  （スライド 1 のナレーション。ノートがあればそのまま、無ければ TODO とスライドの見出し）
  ## 2
  ...

注意: 書き出した台本は必ず人が読み直して整える。記号・数式・略語は読み方を日本語で書き下す
（例: ΔX → デルタ エックス、(n−2)×180° → n 引く 2 かける 180 度）。
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent


def from_pptx(path: Path) -> list[tuple[str, str]]:
    """[(見出し, ノート), ...] をスライド順に返す。"""
    from pptx import Presentation

    def slide_title(slide) -> str:
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                return sh.text_frame.text.strip().splitlines()[0]
        return ""

    prs = Presentation(str(path))
    out = []
    for s in prs.slides:
        notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
        out.append((slide_title(s), notes))
    return out


def from_marp(path: Path) -> list[tuple[str, str]]:
    """Marp の Markdown から [(見出し, ノート), ...] を返す。分割と空スライドの扱いは md2pptx.py と同じ。"""
    sys.path.insert(0, str(HERE.parent))  # scripts/
    import md2pptx

    text = path.read_text(encoding="utf-8")
    _meta, body = md2pptx.split_frontmatter(text)
    out = []
    for s in md2pptx.split_slides(body):
        md, notes, _directives = md2pptx.extract_notes(s)
        blocks = md2pptx.parse_blocks(md)
        if not blocks and not notes:
            continue  # md2pptx.py もこの枚は作らない
        m = re.search(r"^#{1,3}\s+(.+)$", md, flags=re.M)
        title = re.sub(r"\*\*|\$", "", m.group(1)).strip() if m else ""
        out.append((title, "\n\n".join(notes).strip()))
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck", help="slides.pptx または Marp の slides.md")
    ap.add_argument("-o", "--out", default="narration.md")
    args = ap.parse_args()

    deck = Path(args.deck)
    slides = from_marp(deck) if deck.suffix.lower() in (".md", ".markdown") else from_pptx(deck)
    lines = [f"# {deck.stem} ナレーション台本", "",
             "<!-- 各 '## 番号' の下に、そのスライドで話す文章を書く。番号はスライド番号（1 始まり）。 -->", ""]
    n_notes = 0
    for i, (title, notes) in enumerate(slides, 1):
        lines.append(f"## {i}")
        if title:
            lines.append(f"<!-- {title} -->")
        if notes:
            lines.append(notes)
            n_notes += 1
        else:
            lines.append("TODO: このスライドのナレーションを書く。")
        lines.append("")
    Path(args.out).write_text("\n".join(lines), encoding="utf-8")
    print(f"wrote {args.out}: {len(slides)} slides, notes on {n_notes}")


if __name__ == "__main__":
    main()
