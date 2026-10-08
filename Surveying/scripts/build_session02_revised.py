#!/usr/bin/env python3
"""第2回改訂版の図とPPTXを、改訂版Markdownから再生成する。"""
from pathlib import Path
import math
import os
import re
import tempfile
import atexit
import shutil
import sys
sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[1]
# フォント・描画の一時キャッシュも作業対象の配下に限定する。
_cache = tempfile.mkdtemp(prefix='.session02_cache_', dir=ROOT)
atexit.register(shutil.rmtree, _cache, ignore_errors=True)
os.environ['MPLCONFIGDIR'] = _cache
os.environ['XDG_CACHE_HOME'] = _cache
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Circle, FancyBboxPatch
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
import md2pptx as m

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'materials/slides/第02回_測量の基礎と誤差論/改訂版_20261008'
FIG = OUT / 'diagrams'
FIG.mkdir(parents=True, exist_ok=True)
# 中間ファイルもSurveying以下に置く。終了時に自分で作った一時領域だけ削除。
# Noto Sans CJK JP が無い端末（Windows 等）ではメイリオで代替する
# （Windows の Noto Sans JP は可変フォントで、matplotlib では極細になるため使わない）。
plt.rcParams.update({'font.family': ['Noto Sans CJK JP', 'Meiryo'], 'font.size': 14,
                     'axes.spines.top': False, 'axes.spines.right': False,
                     'svg.fonttype': 'path', 'axes.unicode_minus': False})
NAVY, TEAL, ORANGE, GRAY = '#1f4e79', '#13877b', '#bd512b', '#526273'

def save(fig, name):
    svg = FIG / f'{name}.svg'
    fig.savefig(svg, bbox_inches='tight', facecolor='white')
    svg.write_text('\n'.join(line.rstrip() for line in svg.read_text(encoding='utf-8').splitlines())+'\n',encoding='utf-8')
    fig.savefig(FIG / f'{name}.png', dpi=180, bbox_inches='tight', facecolor='white')
    plt.close(fig)

def canvas(h=3):
    fig, ax = plt.subplots(figsize=(12,h))
    ax.set_xlim(0,12); ax.set_ylim(0,h); ax.axis('off')
    return fig, ax

def box(ax,x,y,w,h,text,color=NAVY,fs=16):
    ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.04,rounding_size=0.1',
                 linewidth=1.5, edgecolor=color, facecolor='#f3f7fb'))
    ax.text(x+w/2,y+h/2,text,ha='center',va='center',fontsize=fs,color=color,linespacing=1.6)

def arrow(ax,x1,y1,x2,y2,color=GRAY):
    ax.annotate('',(x2,y2),(x1,y1),arrowprops={'arrowstyle':'-|>','lw':2,'color':color})

fig,ax=canvas(2.4)
box(ax,.1,.55,3.3,1.5,'観測値 ＝ 見える\n1000・1002・1001 mm',fs=14)
arrow(ax,3.55,1.3,4.15,1.3)
box(ax,4.3,.55,3.3,1.5,'平均 ＝ 推定する値\n1001 mm',TEAL)
arrow(ax,7.75,1.3,8.35,1.3)
box(ax,8.5,.55,3.3,1.5,'真値 X ＝ 知りたい値\n正確には分からない',ORANGE)
ax.text(8,.22,'平均と真値が一致するとは限らない',ha='center',fontsize=15,color=ORANGE)
save(fig,'true_value')

fig,axes=plt.subplots(1,3,figsize=(12,3.3))
fig.suptitle('コイン1枚：表なら +1 mm、裏なら −1 mm',fontsize=14,color=NAVY,y=.99)
for ax,k in zip(axes,[2,4,32]):
    total=2*np.arange(k+1)-k
    counts=np.array([math.comb(k,int(j)) for j in range(k+1)])
    p=counts/2**k
    ax.bar(total,p,width=1.6,color=TEAL,alpha=.7)
    ax.set_title(f'コイン {k} 枚',fontsize=16,color=NAVY)
    ax.set_xlabel('合計（mm）',fontsize=12)
    ax.set_ylabel('出る割合',fontsize=12)
    ax.tick_params(labelsize=12)
    if k<32:
        for x,y,c in zip(total,p,counts):
            ax.text(x,y+.018,f'{c}通り',ha='center',fontsize=12,color=NAVY)
        ax.text(.97,.94,f'全 {2**k} 通り',transform=ax.transAxes,ha='right',va='top',fontsize=12,color=GRAY)
        ax.set(xlim=(-k-1.5,k+1.5),ylim=(0,.65),xticks=total)
    else:
        x=np.linspace(-32,32,600)
        # 棒の間隔は2 mmなので、密度を2倍して出る割合と比較する。
        normal=2*np.exp(-x*x/(2*k))/np.sqrt(2*np.pi*k)
        ax.plot(x,normal,color=ORANGE,lw=2,label='なめらかな山（正規分布）')
        ax.set(xlim=(-33,33),ylim=(0,.21),xticks=[-32,-16,0,16,32])
        ax.legend(fontsize=12,loc='upper center',frameon=False,handlelength=1)
fig.tight_layout(rect=(0,0,1,1),pad=.8); save(fig,'clt')

fig,axes=plt.subplots(1,4,figsize=(12,3.3))
# 平均を中心にそろえた固定点群。散らばりと中心のずれを別々に変える。
pts=np.array([[-.13,.02],[.03,.12],[.1,-.07],[-.03,-.12],[.14,.06],[-.1,.14]])
pts=pts-pts.mean(axis=0)
wide=np.array([[-.65,.2],[.1,.68],[.6,-.35],[-.35,-.58],[.55,.4],[-.3,-.35]])
wide=wide-wide.mean(axis=0)
offset=np.array([.52,.3])
sets=[pts,pts+offset,wide,wide+offset]
titles=['精度 高・妥当性 高','精度 高・妥当性 低','精度 低・妥当性 高','精度 低・妥当性 低']
subs=['そろって中心','そろうが中心からずれる','散らばるが中心は合う','散らばり、中心もずれる']
for ax,p,title,sub in zip(axes,sets,titles,subs):
    for rad in [1,.66,.33]: ax.add_patch(Circle((0,0),rad,fill=False,edgecolor=GRAY,alpha=.35,lw=1.5))
    ax.axhline(0,color=GRAY,alpha=.2,lw=.8); ax.axvline(0,color=GRAY,alpha=.2,lw=.8)
    ax.scatter(*p.T,s=38,color=TEAL,zorder=3)
    ax.plot(0,0,'+',color=ORANGE,ms=19,mew=2.5,zorder=4)
    ax.plot(*p.mean(axis=0),'x',color=NAVY,ms=12,mew=2.5,zorder=5)
    ax.set(xlim=(-1.25,1.25),ylim=(-1.2,1.2)); ax.set_aspect('equal');ax.axis('off')
    ax.set_title(title+'\n'+sub,fontsize=12,color=NAVY,linespacing=1.7)
fig.text(.5,.04,'＋ 真値（的の中心）　× 観測値の平均',ha='center',fontsize=13,color=GRAY)
fig.subplots_adjust(left=.02,right=.98,top=.72,bottom=.15,wspace=.12)
save(fig,'targets')

from scipy.stats import t as student_t
fig,ax=plt.subplots(figsize=(12,3.3))
critical=student_t.ppf(.975,df=5)
x=np.sort(np.r_[np.linspace(-4,4,1201),-critical,critical])
density=student_t.pdf(x,df=5)
ax.fill_between(x,0,density,where=np.abs(x)<=critical,color=TEAL,alpha=.18)
ax.fill_between(x,0,density,where=np.abs(x)>=critical,color=ORANGE,alpha=.5)
ax.plot(x,density,color=NAVY,lw=2)
for boundary in [-critical,critical]:
    ax.vlines(boundary,0,student_t.pdf(boundary,df=5),color=ORANGE,lw=2)
ax.text(0,.17,'95%',ha='center',fontsize=23,color=TEAL)
for sign in [-1,1]:
    ax.annotate('2.5%',xy=(sign*3.15,.011),xytext=(sign*3.35,.10),ha='center',fontsize=14,color=ORANGE,
                arrowprops={'arrowstyle':'->','color':ORANGE})
ax.annotate(r'ここより下が 97.5% → $t_{0.975}$',xy=(critical,.045),xytext=(1.15,.30),
            fontsize=13,color=NAVY,arrowprops={'arrowstyle':'->','color':GRAY})
ax.set(xlim=(-4,4),ylim=(0,.42),xlabel=r'$t$',ylabel='密度',
       xticks=[-4,-critical,0,critical,4],xticklabels=['−4','−2.571','0','+2.571','+4'])
ax.set_title('自由度 5 の t 分布',fontsize=16,color=NAVY)
ax.tick_params(labelsize=12)
fig.tight_layout();save(fig,'ci_tails')

fig,axes=plt.subplots(1,2,figsize=(12,3.3))
shifts=[np.array([3,-2,1,-3,2,-1,-2,2]),np.array([1,2,3,2,1,3,2,2])]
for ax,values,title in zip(axes,shifts,['独立：ばらばらの向き','連動：同じ向き']):
    ax.axhline(0,color=GRAY,lw=1)
    for i,value in enumerate(values,1):
        color=TEAL if value<0 else ORANGE
        arrow(ax,i,0,i,value,color)
        ax.text(i,value+(.18 if value>0 else -.18),f'{value:+d}',ha='center',
                va='bottom' if value>0 else 'top',fontsize=12,color=color)
    ax.set(xlim=(.4,8.6),ylim=(-4,4.2),xticks=range(1,9),yticks=[-3,0,3],ylabel='ずれ（mm）')
    ax.set_title(title,fontsize=16,color=NAVY)
    ax.tick_params(labelsize=12)
    total=int(values.sum())
    summary='合計 0 mm → 平均 ≈ 0' if total==0 else f'合計 +{total} mm → 平均 ≠ 0（偏りが残る）'
    ax.text(.5,-.25,summary,transform=ax.transAxes,ha='center',fontsize=14,color=NAVY)
fig.subplots_adjust(left=.065,right=.98,top=.84,bottom=.25,wspace=.24)
save(fig,'cancel')

fig,axes=plt.subplots(1,2,figsize=(12,3.3))
for ax in axes:
    ax.set_aspect('equal');ax.axis('off')
axes[0].plot([0,3],[0,0],color=TEAL,lw=3)
axes[0].plot([3,3],[0,4],color=TEAL,lw=3)
axes[0].plot([0,3],[0,4],color=ORANGE,lw=3)
axes[0].plot([2.65,2.65,3],[0,.35,.35],color=GRAY,lw=1.2)
axes[0].text(1.5,-.55,'3 mm',ha='center',fontsize=14,color=TEAL)
axes[0].text(3.2,2,'4 mm',va='center',fontsize=14,color=TEAL)
axes[0].text(.85,2.5,'5 mm',ha='center',fontsize=14,color=ORANGE,rotation=53.13)
axes[0].set(xlim=(-1.5,6.5),ylim=(-1,5))
axes[0].set_title('独立なずれ：二乗して足す',fontsize=14,color=NAVY,pad=34)
axes[0].text(.5,1.025,r'$\sqrt{3^2+4^2}=5$ mm',transform=axes[0].transAxes,
             ha='center',va='bottom',fontsize=14,color=NAVY)
axes[1].plot([0,3],[1,1],color=TEAL,lw=4)
axes[1].plot([3,7],[1,1],color=ORANGE,lw=4)
for pos in [0,3,7]: axes[1].plot([pos,pos],[.8,1.2],color=NAVY,lw=1.5)
axes[1].text(1.5,1.5,'3 mm',ha='center',fontsize=14,color=TEAL)
axes[1].text(5,1.5,'4 mm',ha='center',fontsize=14,color=ORANGE)
axes[1].annotate('',(7,0),(0,0),arrowprops={'arrowstyle':'|-|','color':NAVY,'lw':1.5})
axes[1].text(3.5,-.7,'3 + 4 = 7 mm',ha='center',fontsize=16,color=NAVY)
axes[1].set(xlim=(-.5,7.5),ylim=(-2,4))
axes[1].set_title('同じ向きにそろうずれ：そのまま足す',fontsize=14,color=NAVY,pad=34)
axes[1].text(.5,1.025,'3 + 4 = 7 mm',transform=axes[1].transAxes,
             ha='center',va='bottom',fontsize=14,color=NAVY)
fig.tight_layout();save(fig,'variance_add')

fig,ax=canvas(2.7)
box(ax,.1,1.25,2.5,1.05,'同じ対象・単位？\n距離の種類もそろえる',fs=14)
box(ax,3.05,1.25,2.5,1.05,'過失・偏りは？\n照合・補正・再測',fs=14)
box(ax,6,1.25,2.5,1.05,'独立性・精度は？\n観測条件を確認',fs=14)
box(ax,9.0,1.45,2.8,.85,'同精度 → 算術平均',TEAL,14)
box(ax,9.0,.15,2.8,.85,'異精度 → 重み付き平均',TEAL,14)
arrow(ax,2.65,1.8,3,1.8);arrow(ax,5.6,1.8,5.95,1.8)
arrow(ax,8.55,1.8,8.95,1.8);arrow(ax,8.7,1.6,8.95,.65)
ax.text(4.3,.65,'そろわない・不明 → 元の記録へ戻って点検',ha='center',color=ORANGE,fontsize=16)
ax.text(4.3,.1,'100 cm と 1.02 m → 1.00 m と 1.02 m → 平均 1.01 m',ha='center',color=NAVY,fontsize=14)
save(fig,'check_flow')

fig,ax=canvas(2.9)
for i,x in enumerate([.05,2.15,4.25,6.35],1):
 box(ax,x,1.6,1.8,.9,f'観測 {i}\n揺らぎ 4 mm',fs=14)
 arrow(ax,x+.9,1.52,x+.9,.88)
 ax.text(x+.9,.52,'× 1/4',ha='center',color=TEAL,fontsize=19)
 if i<4:ax.text(x+1.97,.54,'+',ha='center',fontsize=22,color=NAVY)
arrow(ax,8.3,1.6,8.85,1.6)
box(ax,9.0,1.0,2.75,1.3,'4回の平均\n揺らぎ 2 mm',TEAL,19)
ax.text(4.1,.02,'各矢印の「4分の1」と、全体の「半分」は別',ha='center',fontsize=15,color=ORANGE)
save(fig,'propagation')

fig,axes=plt.subplots(1,2,figsize=(12,3.3),gridspec_kw={'width_ratios':[1.2,1]})
n=np.linspace(1,32,400)
axes[0].axvspan(16,32,color=ORANGE,alpha=.10)
axes[0].plot(n,4/np.sqrt(n),color=TEAL,lw=2.5,label=r'偶然誤差による平均のばらつき $4/\sqrt{n}$ mm')
axes[0].axhline(1,color=ORANGE,lw=2,label=r'定誤差 $b=1$ mm')
axes[0].scatter([16],[1],color=NAVY,s=48,zorder=4)
axes[0].annotate(r'$n=16$',(16,1),xytext=(10,1.8),fontsize=12,color=NAVY,
                 arrowprops={'arrowstyle':'->','color':GRAY})
axes[0].set(xlim=(1,32),ylim=(0,4.3),xlabel=r'観測回数 $n$',ylabel='mm',xticks=[1,4,8,16,24,32])
axes[0].tick_params(labelsize=12);axes[0].grid(alpha=.15)
fig.legend(*axes[0].get_legend_handles_labels(),loc='upper center',bbox_to_anchor=(.5,1),
           ncol=2,fontsize=12,frameon=False)
axes[1].axis('off')
for y,txt in zip([.96,.78,.60,.38],['1回：4 mm → 4回：2 mm','9回：約1.33 mm → 16回：1 mm',
                                   '半分のばらつきには、4倍の回数','定誤差 $b$ は何回平均しても残る']):
    axes[1].text(.03,y,txt,fontsize=14,color=ORANGE if y<.4 else NAVY,transform=axes[1].transAxes)
axes[1].text(.03,.12,'網掛け部分：回数を増やしても\n真値に近づかない',fontsize=14,color=ORANGE,
             transform=axes[1].transAxes,va='top',linespacing=1.5)
fig.tight_layout(rect=(0,0,1,.85));save(fig,'repeats')

fig,axes=plt.subplots(1,2,figsize=(12,2.85))
a=np.linspace(0,1,101);v=a*a+4*(1-a)**2
axes[0].plot(a,v,color=TEAL,lw=2.5)
for x,y,lab in [(0.5,1.25,'半々：1.25'),(.8,.8,'80%：0.80'),(1,1,'第1だけ：1.00')]:
 axes[0].scatter([x],[y],color=ORANGE,s=35,zorder=3)
 axes[0].annotate(lab,(x,y),xytext=(x-.25,y+.65),fontsize=12,arrowprops={'arrowstyle':'-','color':GRAY})
axes[0].set(xlim=(0,1.05),ylim=(.5,4.1),xlabel='第1観測を採用する割合 α',ylabel='統合した値の分散（mm²）')
axes[0].grid(alpha=.15);axes[1].axis('off')
for y,t in zip([.90,.64,.38,.10],['分散 = α² × 1² + (1−α)² × 2²','半々の配分：分散1.25 mm²','80%と20%：分散0.80 mm²','標準偏差は √0.80 ≈ 0.89 mm']):
 axes[1].text(0,y,t,transform=axes[1].transAxes,fontsize=15,color=TEAL if y<.5 else NAVY)
fig.tight_layout();save(fig,'allocation')

class RevisedBuilder(m.Builder):
    def _write_runs(self, par, runs, size, color=None, bold=False):
        # Markdownリンクを編集可能なPowerPointのハイパーリンクにする。
        for text,fmt in runs:
            pos=0
            for match in re.finditer(r'\[([^\]]+)\]\((https?://[^)]+)\)',text):
                super()._write_runs(par,[(text[pos:match.start()],fmt)],size,color,bold)
                before=len(par.runs)
                super()._write_runs(par,[(match[1],fmt)],size,RGBColor(0x13,0x87,0x7b),bold)
                for run in list(par.runs)[before:]: run.hyperlink.address=match[2]
                pos=match.end()
            before=len(par.runs)
            super()._write_runs(par,[(text[pos:],fmt)],size,color,bold)
            if fmt.get('math'):
                for run in list(par.runs)[before:]: run.font.name='DejaVu Sans'

    def add_slide(self,blocks,notes,directives,paginate):
        citations=[b for b in blocks if b['type']=='quote' and b['text'].startswith('出典')]
        blocks=[b for b in blocks if b not in citations]
        pics=[b for b in blocks if b['type']=='para' and re.fullmatch(r'!\[.*?\]\(.*?\.svg\)',b['text'])]
        if not pics:
            super().add_slide(blocks,notes,directives,paginate)
            slide=self.prs.slides[-1]
        else:
            self.page+=1
            slide=self.prs.slides.add_slide(self.blank)
            title=next(b for b in blocks if b['type']=='heading')
            _,tf=self._textbox(slide,m.MARGIN,Inches(.35),m.SLIDE_W-2*m.MARGIN,m.TITLE_H)
            tf.vertical_anchor=MSO_ANCHOR.MIDDLE
            self._write_runs(tf.paragraphs[0],m.inline_to_runs(title['text']),30,m.ACCENT,True)
            ln=slide.shapes.add_shape(1,m.MARGIN,Inches(1.3),m.SLIDE_W-2*m.MARGIN,Inches(.02))
            ln.fill.solid();ln.fill.fore_color.rgb=m.ACCENT;ln.line.fill.background()
            body=[b for b in blocks if b is not title and b not in pics]
            self._place_body(slide,body,m.MARGIN,Inches(1.5),m.SLIDE_W-2*m.MARGIN,Inches(2.1),20)
            from PIL import Image
            png=OUT/re.search(r'\((.*?)\)',pics[0]['text'])[1].replace('.svg','.png')
            with Image.open(png) as im:w,h=im.size
            maxw,maxh=12.0,2.9
            scale=min(maxw/w,maxh/h)
            pw,ph=w*scale,h*scale
            pic=slide.shapes.add_picture(str(png),Inches((13.333-pw)/2),Inches(3.65+(2.9-ph)/2),Inches(pw),Inches(ph))
            pic._element.nvPicPr.cNvPr.set('descr',re.search(r'!\[(.*?)\]',pics[0]['text'])[1])
            self._footer(slide,paginate);self._notes(slide,notes)
        if citations:
            _,tf=self._textbox(slide,m.MARGIN,Inches(6.72),m.SLIDE_W-2*m.MARGIN,Inches(.28))
            self._write_runs(tf.paragraphs[0],[(b['text'],{}) for b in citations],10,m.GRAY)


def build():
    import shutil
    with tempfile.TemporaryDirectory(prefix='.build_session02_',dir=OUT) as scratch:
        tempfile.tempdir=scratch
        text=(OUT/'slides_revised.md').read_text(encoding='utf-8')
        meta,body=m.split_frontmatter(text)
        b=RevisedBuilder('Noto Sans CJK JP','測量学 第2回｜図解・やさしい解説版')
        for s in m.split_slides(body):
            md,notes,directives=m.extract_notes(s)
            b.add_slide(m.parse_blocks(md),notes,directives,True)
        b.save(str(OUT/'測量学_第02回_図解改訂版.pptx'))
        print(f'生成：{b.page} slides')

if __name__=='__main__': build()
