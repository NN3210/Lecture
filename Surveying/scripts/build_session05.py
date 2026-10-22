#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""第5回「トラバース測量 I」の slides.md から、編集可能な pptx を作る。

md2pptx.Builder を継承し、次を加える。
- 図（diagrams/*.svg → 同名 .png、illustrations/*.png）を本文の流れの中に置き、残りの高さに合わせて縮小する。
- ブロック数式は KaTeX → MathML → Office 数式（OMML）に変換し、PowerPoint で編集できる数式にする
  （scripts/katex_mathml.mjs と scripts/omml_equations.py）。
- 表紙（h1）は左に題名、右にイラストの 2 段組にする。
- **強調** と引用は Marp と同じ青緑にする。

使い方（Surveying/ で）:
    python scripts/build_session05.py            # materials/slides/第05回_トラバース測量I/slides.pptx を作る
    python scripts/build_session05.py --no-omml  # 数式を Unicode 文字のまま置く（Office がない環境の検証用）
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import md2pptx as m  # noqa: E402
from omml_equations import patch_pptx  # noqa: E402
from PIL import Image  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

ROOT = HERE.parent
SRC_DIR = ROOT / "materials" / "slides" / "第05回_トラバース測量I"
TEAL = RGBColor(0x11, 0x7E, 0x7A)
IMG_RE = re.compile(r"^!\[([^\]]*)\]\(([^)\s]+)\)$")
LINK_RE = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+)\)")
NODE_CANDIDATES = [
    shutil.which("node"),
    str(Path.home() / ".cache" / "codex-runtimes" / "codex-primary-runtime" / "dependencies" / "node" / "bin" / "node.exe"),
]


class Session05Builder(m.Builder):
    """図と Office 数式に対応したビルダー。"""

    def __init__(self, font: str, title: str, base_dir: Path, use_omml: bool = True):
        super().__init__(font, title)
        self.base_dir = Path(base_dir)
        self.use_omml = use_omml
        self.maths: dict[str, dict] = {}
        self.missing: list[str] = []
        self.warnings: list[str] = []
        self._body_width = m.SLIDE_W - 2 * m.MARGIN

    # --- 文字列 -------------------------------------------------------------
    def _write_runs(self, par, runs, size, color=None, bold=False):
        expanded = []
        for text, fmt in runs:
            if fmt or "](http" not in text:
                expanded.append((text, fmt))
                continue
            pos = 0
            for mm in LINK_RE.finditer(text):
                if mm.start() > pos:
                    expanded.append((text[pos:mm.start()], {}))
                expanded.append((mm.group(1), {"link": mm.group(2)}))
                pos = mm.end()
            if pos < len(text):
                expanded.append((text[pos:], {}))
        for text, fmt in expanded:
            before = len(par.runs)
            super()._write_runs(par, [(text, fmt)], size, color, bold)
            for r in par.runs[before:]:
                if fmt.get("link"):
                    r.hyperlink.address = fmt["link"]
                    r.font.color.rgb = TEAL
                elif fmt.get("bold") and color is None and not bold:
                    r.font.color.rgb = TEAL

    # --- ブロックの下ごしらえ -------------------------------------------------
    @staticmethod
    def _as_image(block: dict):
        if block["type"] != "para":
            return None
        mm = IMG_RE.match(block["text"].strip())
        return {"type": "image", "alt": mm.group(1), "src": mm.group(2)} if mm else None

    def _prepare(self, blocks: list[dict]) -> list[dict]:
        return [self._as_image(b) or b for b in blocks]

    def _image_file(self, src: str):
        p = self.base_dir / src
        if p.suffix.lower() == ".svg":
            p = p.with_suffix(".png")
        return p if p.exists() else None

    def _fit_image(self, block: dict, max_w, max_h):
        """画像を (path, 幅, 高さ) に収める。ファイルがなければ None。"""
        path = self._image_file(block["src"])
        if path is None:
            self.missing.append(block["src"])
            return None
        with Image.open(path) as im:
            wpx, hpx = im.size
        scale = min(max_w / wpx, max_h / hpx)
        return path, int(wpx * scale), int(hpx * scale)

    @staticmethod
    def _math_height(latex: str, eq_size: int):
        lines = [ln for ln in re.split(r"\\\\", latex) if ln.strip()] or [latex]
        h = 0
        for ln in lines:
            tall = any(k in ln for k in ("\\frac", "\\dfrac", "\\sum", "\\sqrt", "\\boxed"))
            h += Pt(eq_size * (2.9 if tall else 2.1))
        return int(h) + Inches(0.2)

    def _table_row_heights(self, rows, width, tsize: int) -> list[int]:
        """各行の高さ。セル内の折り返し行数（列幅は等分）を見積もって決める。"""
        ncols = max(len(r) for r in rows) if rows else 1
        col_pt = (width / 914400) * 72 / ncols - 2 * 0.06 * 72
        cols = max(4.0, col_pt / tsize)
        heights = []
        for r in rows:
            lines = max(self._wrapped(c, cols) for c in r) if r else 1
            heights.append(int(Pt(tsize * (1.35 * lines + 0.65))))
        return heights

    def _add_table(self, slide, rows, left, y, width, size):
        """親クラスの表に、折り返しを見込んだ行高さを付ける。"""
        if not rows:
            return y
        tsize = max(12, size - 4)
        heights = self._table_row_heights(rows, width, tsize)
        y_end = super()._add_table(slide, rows, left, y, width, size)
        table = slide.shapes[-1].table
        for row, h in zip(table.rows, heights):
            row.height = h
        return y + sum(heights) + Inches(0.15)

    def _block_height(self, b: dict, size: int, cols: float):
        if b["type"] == "image":
            return None
        if b["type"] == "table":
            return sum(self._table_row_heights(b["rows"], self._body_width, max(12, size - 4))) + Inches(0.15)
        if b["type"] == "math":
            return self._math_height(b["latex"], size + 2)
        return int(Pt(size * 1.55) * self._estimate_lines([b], cols)) + Inches(0.15)

    # --- 1 枚 ---------------------------------------------------------------
    def add_slide(self, blocks: list[dict], notes: list[str], directives: dict, paginate: bool):
        blocks = self._prepare(blocks)
        title_block = next((b for b in blocks if b["type"] == "heading" and b["level"] <= 2), None)
        if title_block is not None and title_block["level"] == 1:
            self.page += 1
            self._title_slide(blocks, title_block, notes, paginate)
            return
        if not any(b["type"] == "image" for b in blocks):
            super().add_slide(blocks, notes, directives, paginate)
            return
        self.page += 1
        slide = self.prs.slides.add_slide(self.blank)
        body = [b for b in blocks if b is not title_block]
        n_lines = self._estimate_lines([b for b in body if b["type"] != "image"])
        size = 22 if n_lines <= 7 else 20 if n_lines <= 10 else 18 if n_lines <= 13 else 16
        top = self._title(slide, title_block)
        avail_h = m.SLIDE_H - top - m.FOOT_H - Inches(0.1)
        self._place_body(slide, body, m.MARGIN, top, m.SLIDE_W - 2 * m.MARGIN, avail_h, size, center=True)
        self._footer(slide, paginate)
        self._notes(slide, notes)

    def _title(self, slide, title_block):
        if title_block is None:
            return m.MARGIN
        tb, tf = self._textbox(slide, m.MARGIN, Inches(0.35), m.SLIDE_W - 2 * m.MARGIN, m.TITLE_H)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        self._write_runs(p, m.inline_to_runs(title_block["text"]), 30, m.ACCENT, bold=True)
        ln = slide.shapes.add_shape(1, m.MARGIN, Inches(1.3), m.SLIDE_W - 2 * m.MARGIN, Emu(18000))
        ln.fill.solid()
        ln.fill.fore_color.rgb = m.ACCENT
        ln.line.fill.background()
        return Inches(1.5)

    def _title_slide(self, blocks, title_block, notes, paginate):
        slide = self.prs.slides.add_slide(self.blank)
        body = [b for b in blocks if b is not title_block]
        image = next((b for b in body if b["type"] == "image"), None)
        texts = [t for b in body if b["type"] != "image" for t in self._flatten(b)]
        bar = slide.shapes.add_shape(1, 0, 0, m.SLIDE_W, Inches(0.18))
        bar.fill.solid()
        bar.fill.fore_color.rgb = m.ACCENT
        bar.line.fill.background()
        top, box_h = Inches(1.9), Inches(3.6)
        fit = self._fit_image(image, Inches(6.0), box_h) if image else None
        if fit:
            path, w, h = fit
            pic = slide.shapes.add_picture(str(path), m.SLIDE_W - m.MARGIN - w, top + int((box_h - h) / 2), w, h)
            pic._element.nvPicPr.cNvPr.set("descr", image["alt"])
            tb, tf = self._textbox(slide, m.MARGIN, top, Inches(6.1), box_h)
            align = PP_ALIGN.LEFT
        else:
            tb, tf = self._textbox(slide, m.MARGIN, top, m.SLIDE_W - 2 * m.MARGIN, box_h)
            align = PP_ALIGN.CENTER
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = align
        self._write_runs(p, m.inline_to_runs(title_block["text"]), 44, m.ACCENT, bold=True)
        for k, t in enumerate(texts):
            p = tf.add_paragraph()
            p.alignment = align
            p.space_before = Pt(18 if k == 0 else 8)
            self._write_runs(p, m.inline_to_runs(t), 24 if k == 0 else 20, m.GRAY)
        self._footer(slide, paginate)
        self._notes(slide, notes)

    def save(self, path):
        """保存前に全段落へ日本語の禁則処理（句読点のぶら下げ）と言語 ja-JP を付け、行頭の「。」を防ぐ。"""
        A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        for slide in self.prs.slides:
            for p in slide._element.xpath(".//a:p"):
                pPr = p.get_or_add_pPr()
                pPr.set("eaLnBrk", "1")
                pPr.set("hangingPunct", "1")
                for r in p.xpath("./a:r"):
                    rPr = r.get_or_add_rPr()
                    rPr.set("lang", "ja-JP")
                    # 言語を ja-JP にすると和文は a:ea のフォントで描かれるので、Meiryo を明示する
                    latin = rPr.find(A + "latin")
                    if latin is None:
                        latin = rPr.makeelement(A + "latin", {"typeface": self.font})
                        rPr.append(latin)
                    if rPr.find(A + "ea") is None:
                        latin.addnext(rPr.makeelement(A + "ea", {"typeface": self.font}))
        super().save(path)

    def _place_body(self, slide, body, left, top, width, avail_h, size, center=False):
        """ブロックを上から順に置く。画像は残りの高さに収め、数式は Office 数式の仮文字列を置く。"""
        cols = (width / 914400) * 72 / size
        line_h = Pt(size * 1.55)
        self._body_width = width
        heights = [self._block_height(b, size, cols) for b in body]
        fixed = sum(h for h in heights if h is not None)
        n_img = sum(1 for h in heights if h is None)
        images: dict[int, tuple | None] = {}
        if n_img:
            room = max(Inches(1.5), int((avail_h - fixed - Inches(0.12) * n_img) / n_img))
            for i, b in enumerate(body):
                if b["type"] == "image":
                    images[i] = self._fit_image(b, width, room)
                    heights[i] = (images[i][2] + Inches(0.12)) if images[i] else Inches(0.4)
        total = sum(heights)
        y = top
        if center and total < avail_h:
            y += min(Inches(0.3), int((avail_h - total) / 2))
        pending: list[dict] = []

        def flush():
            nonlocal y, pending
            if not pending:
                return
            n = sum(self._estimate_lines([b], cols) for b in pending)
            h = int(line_h * n) + Inches(0.15)
            tb, tf = self._textbox(slide, left, y, width, h)
            first = True

            def para():
                nonlocal first
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                return p

            for b in pending:
                if b["type"] == "heading":
                    p = para()
                    self._write_runs(p, m.inline_to_runs(b["text"]), size + 2, m.ACCENT, bold=True)
                    p.space_before = Pt(6)
                elif b["type"] == "list":
                    for it in b["items"]:
                        p = para()
                        p.level = it["level"]
                        bullet = "• " if it["level"] == 0 else "– "
                        self._write_runs(p, [(bullet, {})] + m.inline_to_runs(it["text"]), size - 2 * it["level"])
                        p.space_after = Pt(4)
                elif b["type"] == "quote":
                    p = para()
                    self._write_runs(p, [("▍", {})] + m.inline_to_runs(b["text"]), size, TEAL)
                    p.space_before = Pt(6)
                elif b["type"] == "code":
                    for ln in b["text"].split("\n"):
                        p = para()
                        self._write_runs(p, [(ln, {"code": True})], size - 4)
                else:
                    p = para()
                    self._write_runs(p, m.inline_to_runs(b["text"]), size)
                    p.space_after = Pt(6)
            y += h
            pending = []

        for i, b in enumerate(body):
            if b["type"] == "table":
                flush()
                y = self._add_table(slide, b["rows"], left, y, width, size)
            elif b["type"] == "math":
                flush()
                y = self._add_equation(slide, b["latex"], left, y, width, size)
            elif b["type"] == "image":
                flush()
                y = self._add_image(slide, images.get(i), b, left, y, width)
            else:
                pending.append(b)
        flush()
        if y > top + avail_h + Inches(0.3):
            self.warnings.append(f"slide {self.page}: 本文がはみ出す可能性（下端 {y / 914400:.2f} in）")

    def _add_image(self, slide, fit, block, left, y, width):
        if fit is None:
            tb, tf = self._textbox(slide, left, y, width, Inches(0.4))
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            self._write_runs(p, [(f"［図：{block['alt']}］", {})], 14, m.GRAY)
            return y + Inches(0.4)
        path, w, h = fit
        pic = slide.shapes.add_picture(str(path), left + int((width - w) / 2), y, w, h)
        pic._element.nvPicPr.cNvPr.set("descr", block["alt"])
        return y + h + Inches(0.12)

    def _add_equation(self, slide, latex, left, y, width, size):
        eq_size = size + 2
        tex = re.sub(r"\s*\n\s*", " ", latex.strip())
        h = self._math_height(tex, eq_size)
        tb, tf = self._textbox(slide, left, y, width, h)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        if self.use_omml:
            eid = f"EQM{len(self.maths) + 1:05d}"
            self.maths[eid] = {"tex": tex, "display": True}
            self._write_runs(p, [(eid, {"math": True})], eq_size)
        else:
            self._write_runs(p, m.split_scripts(m.latex_to_unicode(tex)), eq_size)
        return y + h + Inches(0.05)


def build(md_path: Path, out_path: Path, font: str, use_omml: bool, build_dir: Path | None = None):
    text = md_path.read_text(encoding="utf-8")
    meta, body = m.split_frontmatter(text)
    slides_md = m.split_slides(body)
    paginate = meta.get("paginate", "false").lower() == "true"
    title = md_path.stem
    for b in m.parse_blocks(m.extract_notes(slides_md[0])[0]):
        if b["type"] == "heading":
            title = re.sub(r"\*\*|\$", "", b["text"])
            break
    builder = Session05Builder(font, f"測量学 第5回　{title}", md_path.parent, use_omml)
    for s in slides_md:
        md, notes, directives = m.extract_notes(s)
        blocks = m.parse_blocks(md)
        if not blocks and not notes:
            continue
        builder.add_slide(blocks, notes, directives, paginate)

    build_dir = build_dir or (md_path.parent / "build")
    build_dir.mkdir(parents=True, exist_ok=True)
    if use_omml and builder.maths:
        candidate = build_dir / "candidate.pptx"
        builder.save(candidate)
        src_json = build_dir / "maths_input.json"
        src_json.write_text(json.dumps(builder.maths, ensure_ascii=False, indent=2), encoding="utf-8")
        node = next((n for n in NODE_CANDIDATES if n and Path(n).exists()), None)
        if node is None:
            raise RuntimeError("node が見つかりません（KaTeX の MathML 変換に必要）。")
        subprocess.run([node, str(HERE / "katex_mathml.mjs"), str(src_json), str(build_dir / "maths.json")], check=True)
        maths = json.loads((build_dir / "maths.json").read_text(encoding="utf-8"))
        n = patch_pptx(candidate, out_path, maths, validation=build_dir / "math-validation.json")
        candidate.unlink()
        src_json.unlink()
        print(f"Office 数式: {n} 式")
    else:
        builder.save(out_path)
    print(f"{out_path}: {builder.page} slides")
    for w in builder.warnings:
        print("warning:", w)
    for s in sorted(set(builder.missing)):
        print("missing image:", s)
    return builder


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except AttributeError:
        pass
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--markdown", default=str(SRC_DIR / "slides.md"))
    ap.add_argument("-o", "--output", default=None)
    ap.add_argument("--font", default="Meiryo")
    ap.add_argument("--no-omml", action="store_true", help="数式を Office 数式にせず Unicode 文字で置く")
    ap.add_argument("--build-dir", default=None, help="中間ファイル（maths.json 等）の置き場。既定は Markdown と同じ階層の build/")
    a = ap.parse_args()
    md_path = Path(a.markdown)
    out = Path(a.output) if a.output else md_path.with_suffix(".pptx")
    build(md_path, out, a.font, not a.no_omml, Path(a.build_dir) if a.build_dir else None)


if __name__ == "__main__":
    main()
