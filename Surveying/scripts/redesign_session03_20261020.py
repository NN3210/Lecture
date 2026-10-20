"""第3回81枚の別版。元PPTXのOffice数式を保持し、図解と互換表示を追加する。"""
from pathlib import Path
from copy import deepcopy
from io import BytesIO
import math, json, re, hashlib
from zipfile import ZipFile
from lxml import etree as E
import fitz
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'materials/slides/第03回_測量の座標'
OUT=SRC
REVIEW=ROOT/'notes/session03_visual_20261020'
REVIEW.mkdir(parents=True,exist_ok=True)
NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main','mc':'http://schemas.openxmlformats.org/markup-compatibility/2006','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
U=9525
BG='FAF9F5'; NAVY='173C4D'; TEAL='117E7A'; BLUE='287EAF'; GOLD='B57B16'; RED='C65C42'; INK='243F4C'; GRAY='647880'; PALE='E8F2EF'; WHITE='FFFFFF'; LINE='D6E2E2'
FONT='Noto Sans CJK JP'
p=Presentation(SRC/'slides.pptx'); pdf=fitz.open(SRC/'slides.pdf')
hashes={f:hashlib.sha256((SRC/f).read_bytes()).hexdigest() for f in ['slides.md','slides.pptx','slides.pdf']}

def px(x): return round(x*U)
def rgb(c):return RGBColor.from_string(c)
def rect(s,x,y,w,h,c=WHITE,line=None,r=0):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,px(x),px(y),px(w),px(h))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(c)
    if line:sh.line.color.rgb=rgb(line);sh.line.width=Pt(1)
    else:sh.line.fill.background()
    if r: sh.adjustments[0]=.08
    return sh
def text(s,t,x,y,w,h,size=28,c=INK,bold=False,align='left'):
    sh=s.shapes.add_textbox(px(x),px(y),px(w),px(h));tf=sh.text_frame
    tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for j,ln in enumerate(t.split('\n')):
        q=tf.paragraphs[0] if j==0 else tf.add_paragraph();q.text=ln
        q.font.name=FONT;q.font.size=Pt(size*.75);q.font.bold=bold;q.font.color.rgb=rgb(c)
        q.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
        q.space_after=Pt(5);q.line_spacing=1.15
    return sh
def line(s,x,y,X,Y,c=GRAY,w=2,arrow=False,dash=False):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,px(x),px(y),px(X),px(Y));sh.line.color.rgb=rgb(c);sh.line.width=Pt(w*.75)
    if arrow:
        a=OxmlElement('a:tailEnd');a.set('type','triangle');a.set('w','med');a.set('len','med');sh._element.spPr.get_or_add_ln().append(a)
    if dash:
        a=OxmlElement('a:prstDash');a.set('val','dash');sh._element.spPr.get_or_add_ln().append(a)
    return sh
def ellipse(s,x,y,w,h,c=TEAL,fill=True):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,px(x),px(y),px(w),px(h))
    if fill:sh.fill.solid();sh.fill.fore_color.rgb=rgb(c);sh.line.fill.background()
    else:sh.fill.background();sh.line.color.rgb=rgb(c);sh.line.width=Pt(2)
    return sh
def dot(s,x,y,c=TEAL,r=6): return ellipse(s,x-r,y-r,r*2,r*2,c)
def pill(s,t,x,y,w,c=TEAL):rect(s,x,y,w,34,c,r=1);text(s,t,x+10,y+4,w-20,27,18,WHITE,True,'center')
def card(s,k,title,body,x,y,w,h,c=TEAL):
    rect(s,x,y,w,h,WHITE,LINE,1);pill(s,k,x+22,y+22,72,c)
    text(s,title,x+22,y+80,w-44,70,31,c,True);text(s,body,x+22,y+160,w-44,h-172,25)
def take(s,t):
    rect(s,62,629,1156,57,PALE,r=1);rect(s,62,629,6,57,TEAL)
    text(s,t,83,642,1110,34,23,TEAL,True)
def section(n):
    for end,name,c in [(13,'START  |  前回の確認',GRAY),(20,'01  |  場所を伝える',TEAL),(31,'02  |  高さをそろえる',BLUE),(46,'03  |  地球を地図にする',TEAL),(58,'04  |  現地と図面をつなぐ',GOLD),(71,'05  |  角と距離で点を探す',RED),(73,'06  |  現場へ',TEAL),(77,'TRY  |  自分で解く',RED),(81,'GOAL  |  振り返り',TEAL)]:
        if n<=end:return name,c
def header(s,n,title,sub=None):
    sec,c=section(n);rect(s,0,0,1280,8,c)
    text(s,'測量学 第3回  /  '+sec,64,26,1110,27,17,c,True)
    text(s,title,62,69,1156,66,38 if len(title)<31 else 33,NAVY,True)
    if sub:text(s,sub,66,143,1148,50,25,GRAY)
    text(s,f'{n:02d} / 81',1115,691,104,22,14,GRAY,False,'right')
def clean(s,n,title,sub=None):
    for el in list(s.shapes._spTree):
        if E.QName(el).localname not in ['nvGrpSpPr','grpSpPr']:s.shapes._spTree.remove(el)
    s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
    header(s,n,title,sub)
def pic(s,path,x,y,w,h):
    from PIL import Image
    im=Image.open(path);rw,rh=im.size;scale=min(w/rw,h/rh);ww,hh=rw*scale,rh*scale
    s.shapes.add_picture(str(path),px(x+(w-ww)/2),px(y+(h-hh)/2),px(ww),px(hh))
def poly(s,pts,c=TEAL,w=3):
    for a,b in zip(pts,pts[1:]):line(s,*a,*b,c,w)
def globe(s,x,y,d=220):
    ellipse(s,x,y,d,d,'E4F1F2');ellipse(s,x,y,d,d,BLUE,False)
    ellipse(s,x+d*.3,y,d*.4,d,BLUE,False);ellipse(s,x,y+d*.32,d,d*.36,BLUE,False)
    line(s,x+d/2,y-12,x+d/2,y+d+12,GRAY,1);line(s,x,y+d/2,x+d,y+d/2,BLUE,1)
def grid(s,x,y,w,h,step=50):
    for a in range(0,int(w)+1,step):line(s,x+a,y,x+a,y+h,LINE,1)
    for a in range(0,int(h)+1,step):line(s,x,y+a,x+w,y+a,LINE,1)
def north(s,x,y,span=130):
    line(s,x-span,y,x+span,y,GRAY,1.5,True);line(s,x,y+span,x,y-span,TEAL,2,True)
    text(s,'X 北',x-25,y-span-35,75,30,20,TEAL,True)
    text(s,'Y 東',x+span-25,y+10,80,30,20,BLUE,True)

# 原図PDFをOffice数式非対応ビューアのFallbackにのみ埋め込む。
# Choice内の編集可能なOMMLは一文字も変更しない。
math_count=0;fallback_count=0
for n,s in enumerate(p.slides,1):
    for ac in list(s._element.findall('.//mc:AlternateContent',NS)):
        ch=ac.find('mc:Choice',NS);fb=ac.find('mc:Fallback',NS)
        if ch is None or fb is None:continue
        math_count+=len(ch.findall('.//m:oMath',NS))
        xf=ch.find('.//a:xfrm',NS)
        if xf is None:xf=ch.find('.//p:xfrm',NS)
        if xf is None:raise RuntimeError(f'No bounds: {n}')
        off=xf.find('a:off',NS);ext=xf.find('a:ext',NS)
        x,y,w,h=[int(v)/U for v in [off.get('x'),off.get('y'),ext.get('cx'),ext.get('cy')]]
        clip=fitz.Rect(x*.75,y*.75,(x+w)*.75,(y+h)*.75)
        pix=pdf[n-1].get_pixmap(matrix=fitz.Matrix(2.8,2.8),clip=clip,alpha=False)
        im=s.shapes.add_picture(BytesIO(pix.tobytes('png')),px(x),px(y),px(w),px(h))
        el=im._element;el.getparent().remove(el)
        for old in list(fb):fb.remove(old)
        fb.append(el);fallback_count+=1

titles={2:'何回も測ったら、どれだけ確かになる？',3:'信頼できる観測に、大きな重みを',4:'今日のゴール：地図の数字が読める！',7:'小テスト① どの説明が間違い？',8:'小テスト② 5回測った距離の代表値は？',9:'小テスト③ 平均すると誤差はどうなる？',10:'小テスト④ どちらの観測を信頼する？',11:'小テスト⑤「±」の表示を読み解こう',12:'答え合わせ① なぜ、その答えになる？',13:'答え合わせ② 重みと精度を読み解く',15:'点はそのまま、ものさしの向きを変える',16:'今度は、点そのものを回してみよう',20:'同じ杭なのに、数字が違うのはなぜ？',21:'地球は、ほんの少し「横長」',25:'ジオイド高：日本では約20〜45 m',27:'高さは「どこから測るか」で変わる',28:'GNSSで85 m。標高も85 m？',29:'ジオイド高を、となりの点へつなぐ',30:'水はA・B、どちらへ流れる？',31:'30秒チャレンジ：高さの基準',32:'地球に、どんなものさしを当てる？',33:'WGS84・ITRF・日本の基準を整理',34:'緯度・経度・高さで、点を見つける',35:'地球の中心から、3本の軸を伸ばす',36:'基線ベクトル＝「AからBへの移動」',39:'地図で大きく見える＝面積も大きい？',41:'世界を6°ずつに分けると、UTM',42:'UTMと日本の座標系、何が違う？',43:'みんなが使う「位置と高さの出発点」',44:'標高の数字も、更新される',46:'似た言葉は「基準」で見分ける',47:'日本は19の座標系でカバーする',48:'注意！ 測量では「Xが北・Yが東」',49:'北・南・東・西を、＋と−で表す',50:'帯広の住所は「第XIII系」',51:'地図の距離は、場所で少し伸び縮み',53:'まず水平に。次に、基準面へ',54:'最後に「地図の伸び縮み」を直す',55:'1 kmを図面にすると、何cm変わる？',56:'逆向きに計算！ 図面から現地へ',57:'30秒チャレンジ：方角と距離',58:'迷ったら、換算の出発点と到着点を書く',59:'「北」は1つではない！',60:'北のずれを、角度から引く',61:'「北から東へ30°」を数字にする',62:'ななめの移動を「北と東」に分ける',63:'北から120°。点はどちらへ進む？',64:'座標の差から、距離と向きを戻す',65:'電卓の答えだけで、向きは決まらない',66:'答えを逆算して、出発点へ戻ろう',67:'60°なら北東。符号を予想して計算',68:'198°なら南西。マイナスになるのは？',69:'成果表の数字を、地図の上で点検',70:'30秒チャレンジ：向きと角度',71:'座標計算は「式＋方角」でチェック',72:'衛星の測位が、現地の座標になるまで',73:'TSは「角と距離」で点を探す',74:'実践① 座標から探す・座標を求める',75:'実践② 現地の距離を図面へ',76:'計算が止まったら、この順で点検',78:'今日の道具箱：式と意味をセットに',79:'次回までに、この5つを説明しよう',81:'出典・図の扱い'}

# 全81枚共通：章の色、問いを先に置く見出し、ひとこと結論。
for n,s in enumerate(p.slides,1):
    if n==1:continue
    old=list(s.shapes._spTree)
    for el in old:
        if E.QName(el).localname=='sp':
            xf=el.find('.//a:xfrm',NS)
            if xf is None:continue
            y=int(xf.find('a:off',NS).get('y'))/U
            if y<135 or y>675:s.shapes._spTree.remove(el)
    original_title=''
    with ZipFile(SRC/'slides.pptx') as z:
        root=E.fromstring(z.read(f'ppt/slides/slide{n}.xml'))
        for el in root.findall('.//p:sp',NS):
            xf=el.find('.//a:xfrm',NS)
            if xf is not None and int(xf.find('a:off',NS).get('y'))/U==68:
                original_title=''.join(el.xpath('.//a:t/text()',namespaces=NS));break
    header(s,n,titles.get(n,original_title))
    # 結論を元の位置のまま囲む。数式のFallbackの背景色に合わせる。
    has_take=False
    for el in old:
        xf=el.find('.//a:xfrm',NS)
        if xf is not None and xf.find('a:off',NS) is not None and abs(int(xf.find('a:off',NS).get('y'))/U-632)<1:has_take=True
    if has_take:
        sh=rect(s,57,627,1166,62,BG,LINE,1);s.shapes._spTree.remove(sh._element);s.shapes._spTree.insert(2,sh._element)
        rect(s,57,627,5,62,TEAL)

redrawn=[]
def redraw(n,title,sub=None):
    s=p.slides[n-1]
    if s._element.findall('.//m:oMath',NS):raise ValueError(f'Would discard mathematics on {n}')
    clean(s,n,title,sub);redrawn.append(n);return s

# 01 表紙：大きな問いと、現場のイラスト。
s=redraw(1,'地図の「ここ！」を、数字にしよう')
pic(s,SRC/'illustrations/title_hero.png',520,140,730,470)
pill(s,'測量の座標',65,170,190)
text(s,'待ち合わせ場所も、\n畑の高さも、\n数字で伝えられる。',65,236,510,230,44,NAVY,True)
text(s,'座標 → 地球と高さ → 地図 → 座標計算',66,540,1030,50,28,TEAL,True)
take(s,'今日の合言葉：まず図で予想。次に式で計算。最後に符号で確認。')
text(s,'背景：AI生成の概念イラスト',920,604,290,23,15,GRAY,False,'right')

s=redraw(5,'90分の冒険マップ','いまどこを学んでいるか、上の章ラベルで確かめよう。')
stages=[('01','前回を思い出す','0〜16分\n小テストと解説',GRAY),('02','地球と高さ','16〜34分\n座標・ジオイド',BLUE),('03','地図へ写す','34〜56分\n測地系・投影・距離',TEAL),('04','点を探す','56〜68分\n方向角・計算・機器',RED),('05','自分で解く','68〜86分\n手計算テスト',GOLD),('06','振り返る','86〜90分\nまとめ・提出',TEAL)]
for i,(k,t,b,c) in enumerate(stages):
    x=65+i*194;rect(s,x,250,178,300,WHITE,LINE,1);pill(s,k,x+20,275,65,c);text(s,t,x+15,345,155,58,25,c,True);text(s,b,x+15,425,155,100,21)
    if i<5:line(s,x+180,395,x+192,395,c,2,True)
take(s,'難しい名前も「何をそろえる約束？」と考えると、つながって見える。')

s=redraw(6,'まずは10分。前回の知識を起こそう')
ellipse(s,120,205,330,330,PALE);text(s,'10',130,260,310,130,110,TEAL,True,'center');text(s,'分',230,405,110,65,43,TEAL,True,'center')
text(s,'5問 × 2点 ＝ 10点',535,230,630,60,42,NAVY,True)
text(s,'五肢択一・電卓なし\n式を書いて、概算で選ぼう。\n答えの理由も、ひとこと考える。',535,335,630,185,30)
take(s,'平均・ばらつき・重み・機器の精度。前回の道具を使ってみよう。')

s=redraw(14,'「校庭のここ！」を、言葉だけで伝えられる？','原点・向き・単位を決めると、同じ場所を共有できる。')
grid(s,100,230,550,300,50);line(s,100,530,690,530,BLUE,3,True);line(s,100,530,100,210,TEAL,3,True)
dot(s,400,330,RED,10);text(s,'待ち合わせの木',420,290,240,45,28,RED,True)
line(s,400,330,400,530,RED,2,False,True);line(s,100,330,400,330,RED,2,False,True)
text(s,'原点：校門',80,550,230,40,24);text(s,'東へ30 m',260,550,220,40,25,BLUE,True);text(s,'北へ20 m',115,285,220,40,25,TEAL,True)
card(s,'約束','数字の住所をつくる','① 原点＝校門\n② 軸＝北と東\n③ 単位＝m',770,215,430,360)
take(s,'測量の平面直角座標なら、この点は（X, Y）＝（20, 30）m。')

s=redraw(17,'みかんの皮を、破らず平らにできる？','地球は曲面。紙の地図は平面。この違いが「ひずみ」を生む。')
ellipse(s,130,230,300,300,'EAA940');ellipse(s,160,250,240,260,'F4BE5B',False)
for offset in [-70,0,70]:ellipse(s,250+offset*.3,230,60+abs(offset),300,'BC7A20',False)
text(s,'丸い地球',150,550,260,45,29,NAVY,True,'center');line(s,490,370,690,370,TEAL,6,True)
for i in range(4):
    sh=s.shapes.add_shape(MSO_SHAPE.DIAMOND,px(755+i*100),px(240),px(88),px(285));sh.fill.solid();sh.fill.fore_color.rgb=rgb('F4BE5B');sh.line.fill.background()
text(s,'すき間を埋めるには\n伸ばす・縮める必要がある',740,540,450,70,26,RED,True,'center')
take(s,'角・面積・距離を、広い範囲で全部同時に正しく保つことはできない。')

s=redraw(18,'座標にも「送り状」が必要だ','数字だけを渡しても、同じ場所にたどり着けるとは限らない。')
card(s,'01','地球の基準と時点','測地系は？\nいつの座標？\n成果の元期を記録する。',65,218,365,360,TEAL)
card(s,'02','軸・単位・系番号','緯度経度？ 平面座標？\nXとYはどちら向き？\nm？ 度？ 何系？',457,218,365,360,BLUE)
card(s,'03','高さの基準','楕円体高？ 標高？\n成果とジオイドモデルも\n一緒に伝える。',849,218,365,360,GOLD)
take(s,'CRS（座標参照系）は、座標のものさしを地球の基準に結び付ける約束。')

s=redraw(19,'同じ点を、4種類の「住所」で呼ぶ','目的が変わると、使いやすい座標も変わる。点そのものは変わらない。')
labels=[('地心直交座標','地球の中心から\nX・Y・Zで表す','GNSSの計算',BLUE),('地理座標','緯度・経度・\n楕円体高で表す','地球上の位置',TEAL),('投影座標＋標高','平面のX・Yと\n標高Hで表す','地図・設計',GOLD),('局所座標','現場の原点から\n使いやすく表す','施工・実習',RED)]
for i,(a,b,c,col) in enumerate(labels):
    x=65+i*294;rect(s,x,230,270,340,WHITE,LINE,1);text(s,a,x+17,258,236,53,28,col,True);text(s,b,x+17,340,236,110,28);pill(s,c,x+17,500,236,col)
    if i<3:line(s,x+272,390,x+290,390,GRAY,2,True)
take(s,'「どの座標？」を先に確認。変換すると数字は変わるが、場所は同じ。')

s=redraw(22,'コップの水は、何に対して水平？','静止した水面は、その場所の重力方向に垂直になる。')
rect(s,150,315,350,205,'D6EDF4');line(s,150,250,150,520,NAVY,5);line(s,150,520,500,520,NAVY,5);line(s,500,520,500,250,NAVY,5);line(s,150,315,500,315,BLUE,5)
line(s,325,365,325,580,RED,5,True);text(s,'重力の向き',350,535,230,45,27,RED,True);text(s,'静かな水面',190,260,240,45,29,BLUE,True)
line(s,325,315,355,315,GRAY,2);line(s,355,315,355,345,GRAY,2);line(s,355,345,325,345,GRAY,2)
card(s,'考える','地球全体では？','周りの質量分布によって\n重力の向きが少し変わる。\n\nだから水準面は、滑らかな\n楕円体とぴったり重ならない。',680,210,510,380,BLUE)
take(s,'「水平」は、紙の横方向ではなく、その場所の重力で決まる。')

s=redraw(23,'海の水面を、陸の下まで延ばしたら？','平均海面に近い水準面を延ばしたものが「ジオイド」。')
pts=[(80,380),(230,380),(340,330),(440,240),(545,350),(640,300),(780,360),(1000,380),(1190,380)]
poly(s,pts,'9F8660',6);line(s,80,405,1190,405,BLUE,5);line(s,80,520,1190,520,TEAL,3,False,True)
text(s,'海',115,330,100,44,29,BLUE,True);text(s,'地表',440,203,130,44,29,'876D49',True)
text(s,'ジオイド：標高の基準',720,421,470,44,30,BLUE,True);text(s,'楕円体：計算の基準',720,540,470,44,28,TEAL,True)
pill(s,'地形の凹凸とは別の面',75,220,330,BLUE)
take(s,'正確には「重力の等ポテンシャル面」。水が静止できる基準の面と考えよう。')
text(s,'模式図。面の間隔や地形は誇張。',80,585,800,27,17,GRAY)

s=redraw(24,'潮が満ち引きするのに、標高の0 mは？','長い期間の潮位を観測し、平均的な海面を基準にする。')
rect(s,65,220,740,355,WHITE,LINE,1)
pts=[(90+i*2.3,390-80*math.sin(i/18)-22*math.sin(i/7)) for i in range(300)];poly(s,pts,BLUE,3)
line(s,90,390,780,390,TEAL,3,False,True);text(s,'時間 →',650,530,120,35,23,GRAY);pill(s,'平均海面',100,250,150,TEAL)
text(s,'東京湾平均海面',855,255,370,50,32,TEAL,True);text(s,'T.P.',855,315,330,65,50,NAVY,True)
text(s,'日本水準原点の標高',855,425,350,40,24);text(s,'24.3900 m',855,474,350,60,41,NAVY,True)
take(s,'その瞬間の海面を「0 m」にしているわけではない。')
text(s,'波形は説明用の模式図（実測値ではありません）。',85,584,920,28,17,GRAY)

s=redraw(26,'地面・水の基準・計算の基準を重ねる','同じ点でも、どの面から測るかで高さの数字が変わる。')
poly(s,[(80,325),(210,320),(330,260),(440,290),(540,235),(670,310),(820,275)],'9F8660',5)
poly(s,[(80,430),(240,420),(420,432),(620,416),(820,427)],BLUE,5)
line(s,80,535,820,535,TEAL,4)
dot(s,540,235,RED,9);text(s,'P',560,214,50,45,30,RED,True)
for x,y,Y,c,label in [(580,235,535,TEAL,'h'),(680,310,418,BLUE,'H'),(760,425,535,GOLD,'N')]:
    line(s,x,y,x,Y,c,3,True);text(s,label,x+12,(y+Y)/2-20,50,45,31,c,True)
# Hの上端をPの高さに合わせる（各面の模式的な高さの差を示す）。
line(s,680,235,680,310,BLUE,3);line(s,540,235,690,235,GRAY,1,False,True)
text(s,'地表',880,250,300,45,32,'876D49',True);text(s,'実際の地面',880,300,300,40,24)
text(s,'ジオイド',880,390,300,45,32,BLUE,True);text(s,'標高 H の基準',880,440,300,40,24)
text(s,'楕円体',880,510,300,45,32,TEAL,True);text(s,'楕円体高 h の基準',880,558,340,40,24)
take(s,'h ＝ H ＋ N。全体の高さ ＝ 標高 ＋ 基準面どうしの差。')

s=redraw(37,'地球を紙に写す、3つの作戦','これは投影面のイメージ。実際の変換は数学的な式で行う。')
for i,(title,caption,c) in enumerate([('円筒','円筒を開く',BLUE),('円錐','円錐を開く',TEAL),('平面','平面に写す',GOLD)]):
    x=70+i*395;rect(s,x,215,355,385,WHITE,LINE,1);text(s,title,x+25,240,305,48,33,c,True,'center');globe(s,x+92,335,170)
    if i==0:
        line(s,x+62,317,x+62,524,c,3);line(s,x+292,317,x+292,524,c,3);ellipse(s,x+62,295,230,50,c,False);ellipse(s,x+62,500,230,50,c,False)
    elif i==1:
        line(s,x+45,530,x+177,305,c,3);line(s,x+177,305,x+310,530,c,3);ellipse(s,x+45,505,265,50,c,False)
    else:line(s,x+35,335,x+320,335,c,5)
    text(s,caption,x+25,553,305,36,25,c,True,'center')
take(s,'地図の目的に合わせて、何を正しく保つかを選ぶ。')

s=redraw(38,'地球の「しま模様」を、方眼の地図へ','メルカトル図法では、経線は平行、緯線の間隔は高緯度ほど広がる。')
globe(s,95,260,280);line(s,420,405,650,405,TEAL,5,True)
rect(s,710,220,460,370,WHITE,LINE)
for x in range(710,1171,92):line(s,x,220,x,590,BLUE,2)
for y in [220,322,377,405,433,488,590]:line(s,710,y,1170,y,BLUE,2)
text(s,'赤道',1177,390,80,35,21,TEAL,True);text(s,'北ほど間隔が広い',715,174,460,40,26,BLUE,True)
text(s,'南ほど間隔が広い',715,594,460,30,22,BLUE,True)
take(s,'角は正しく保たれる一方、高緯度ほど面積が大きく見える。')

s=redraw(40,'円筒を横にすると、南北に細長い地域に便利','横メルカトル投影（ガウス・クリューゲル投影）へ。')
for x,title,c in [(80,'正軸：赤道を中心に',BLUE),(690,'横軸：中央子午線を中心に',TEAL)]:
    rect(s,x,230,510,340,WHITE,LINE,1);text(s,title,x+18,250,475,45,28,c,True);globe(s,x+150,330,190)
    if x==80:
        ellipse(s,x+125,310,240,45,c,False);ellipse(s,x+125,510,240,45,c,False);line(s,x+125,332,x+125,533,c,3);line(s,x+365,332,x+365,533,c,3)
    else:
        ellipse(s,x+90,310,50,225,c,False);ellipse(s,x+355,310,50,225,c,False);line(s,x+115,310,x+380,310,c,3);line(s,x+115,535,x+380,535,c,3)
line(s,607,410,675,410,GOLD,4,True)
take(s,'中央子午線から東西に離れるほどひずみが増えるので、帯に分けて使う。')

s=redraw(45,'地面は動く。座標の「日付」もそろえよう','セミ・ダイナミック補正：基準日と観測日のずれを扱う。')
card(s,'01','元期 → 今期','既知点の座標に\n補正量を加え、\n観測日の座標へ。',65,230,350,350,BLUE)
card(s,'02','今期で計算','同じ時点の座標で\n観測を計算する。\n（網平均）',465,230,350,350,TEAL)
card(s,'03','今期 → 元期','新点の座標から\n補正量を引き、\n成果の基準日へ。',865,230,350,350,GOLD)
line(s,425,400,453,400,TEAL,4,True);line(s,825,400,853,400,TEAL,4,True)
take(s,'「同じ場所」でも、時点が違えば座標が違う。長い基線ほど影響に注意。')

s=redraw(52,'測った距離を、そのまま地図に使える？','現場から図面まで、4つの距離を順にたどる。')
labels=[('L','斜距離','斜面に沿う見通し',RED),('S','水平距離','水平に直す',BLUE),('S₀','基準面距離','基準面へ下ろす',TEAL),('s','平面距離','投影の縮尺を反映',GOLD)]
for i,(a,b,c,col) in enumerate(labels):
    x=65+294*i;y=210+34*i;rect(s,x,y,266,285,WHITE,LINE,1);text(s,a,x+20,y+20,226,85,63,col,True,'center');text(s,b,x+18,y+124,230,45,30,col,True,'center');text(s,c,x+15,y+200,236,60,23,INK,False,'center')
    if i<3:line(s,x+269,y+142,x+290,y+176,GRAY,3,True)
take(s,'座標計算で使うのは最後の s。斜め → 水平 → 基準面 → 平面。')

s=redraw(77,'提出前に、3つだけチェック')
card(s,'01','まとめる','小テストと\n手計算テスト2ページを\nまとめる。',65,220,365,365,BLUE)
card(s,'02','考えた跡を残す','式・途中計算・単位。\n途中まででも、\n解いたところを提出。',457,220,365,365,TEAL)
card(s,'03','番号を確認','提出用番号は\n配布紙にだけ記入。\n最後に見直そう。',849,220,365,365,GOLD)
take(s,'解説は次回の冒頭。順計算と逆計算を復習しておこう。')

s=redraw(80,'次回は「角」を、もっと正確に測る','今日の座標計算を、現場で使える観測につなげよう。')
north(s,300,410,150);line(s,300,410,470,280,RED,5,True);dot(s,300,410,TEAL,9)
text(s,'見通す向きを\n正確な数字へ。',105,545,450,70,29,TEAL,True,'center')
text(s,'前回演習の解説・第3回の小テスト\n単測法・倍角法・方向法と点検\n偏心観測・TSの自動視準と二軸補正\n巻尺の補正・光波測距の原理',645,255,565,290,29)
take(s,'関数電卓を持参。方向角と順計算・逆計算をもう一度。')

# 数式の例題6枚は、式を左、同じ例の幾何を右に並べる。
side_diagrams=[16,36,63,66,67,68]
for n in side_diagrams:
    s=p.slides[n-1];acs=s._element.findall('.//mc:AlternateContent',NS)
    for ac in acs:
        ch=ac.find('mc:Choice',NS);xf=ch.find('.//a:xfrm',NS)
        if xf is None:continue
        off=xf.find('a:off',NS);ext=xf.find('a:ext',NS);y=int(off.get('y'))/U
        # display数式だけ幅を狭める。段落や条件は全幅を維持。
        ps=ch.findall('.//a:p',NS)
        is_display=any(q.get('alg')=='ctr' or q.get('algn')=='ctr' for q in ch.findall('.//a:pPr',NS))
        if y>=620 or not is_display:continue
        old_w=int(ext.get('cx'))/U;old_h=int(ext.get('cy'))/U;old_x=int(off.get('x'))/U
        ext.set('cx',str(px(745)))
        # PDF内の余白を除いた数式を、同じ実寸で左欄に配置する。
        clip=fitz.Rect(old_x*.75,y*.75,(old_x+old_w)*.75,(y+old_h)*.75)
        spans=[]
        for b in pdf[n-1].get_text('dict',clip=clip)['blocks']:
            for ln in b.get('lines',[]):
                for sp in ln['spans']:
                    if sp['text'].strip():spans.append(fitz.Rect(sp['bbox']))
        if spans:
            tight=spans[0]
            for b in spans[1:]:tight|=b
            tight=fitz.Rect(tight.x0-3,max(clip.y0,tight.y0-4),tight.x1+3,min(clip.y1,tight.y1+5))
        else:tight=clip
        pix=pdf[n-1].get_pixmap(matrix=fitz.Matrix(3,3),clip=tight,alpha=False)
        ww,hh=tight.width/.75,tight.height/.75;fac=min(1,730/ww);ww*=fac;hh*=fac
        sh=s.shapes.add_picture(BytesIO(pix.tobytes('png')),px(68+(745-ww)/2),px(y+(old_h-hh)/2),px(ww),px(hh))
        el=sh._element;el.getparent().remove(el);fb=ac.find('mc:Fallback',NS)
        for child in list(fb):fb.remove(child)
        fb.append(el)
    panel_y=445 if n==67 else (285 if n!=68 else 350);panel_h=610-panel_y
    rect(s,855,panel_y,360,panel_h,WHITE,LINE,1)
    if n==16:
        text(s,'同じ半径の円を動く',875,300,320,40,24,TEAL,True)
        ox,oy=900,550;line(s,ox,oy,1185,oy,GRAY,2,True);line(s,ox,oy,ox,355,GRAY,2,True)
        a=(ox+1.08*165,oy-.91*165);b=(ox+.4803*165,oy-1.328*165)
        # 図をパネル内に収めるため同じ倍率で縮小。
        a=(ox+1.08*145,oy-.91*145);b=(ox+.4803*145,oy-1.328*145)
        line(s,ox,oy,*a,BLUE,3,True);line(s,ox,oy,*b,RED,3,True);dot(s,*a,BLUE);dot(s,*b,RED)
        text(s,'P',a[0]+10,a[1],40,30,23,BLUE,True);text(s,'P′',b[0]+10,b[1]-22,45,30,23,RED,True)
        arc=[(ox+70*math.cos(t),oy-70*math.sin(t)) for t in [math.atan2(.91,1.08)+j*math.pi/180 for j in range(31)]];poly(s,arc,TEAL,3)
        text(s,'30°',980,458,80,40,25,TEAL,True);text(s,'x',1180,550,30,30,20);text(s,'y',878,348,30,30,20)
    elif n==36:
        text(s,'3方向の移動を足す',875,305,320,40,25,TEAL,True)
        a=(900,530);b=(1160,390);mid=(1070,530)
        line(s,*a,*mid,BLUE,3,True);line(s,*mid,1130,480,TEAL,3,True);line(s,1130,480,*b,GOLD,3,True);line(s,*a,*b,RED,4,True)
        dot(s,*a,TEAL);dot(s,*b,RED);text(s,'A',875,540,45,35,25,TEAL,True);text(s,'B',1165,355,40,40,25,RED,True)
        text(s,'ΔX',960,542,70,32,23,BLUE,True);text(s,'ΔY',1050,478,70,32,23,TEAL,True);text(s,'ΔZ',1140,443,70,32,23,GOLD,True)
        text(s,'模式図・縮尺は任意',875,575,320,30,17,GRAY)
    else:
        angle={63:120,66:120,67:60,68:198}[n]
        oy=520 if n==67 else (435 if n!=68 else 450);ox=1035
        north(s,ox,oy,43 if n==67 else (112 if n!=68 else 80))
        radius=50 if n==67 else (112 if n!=68 else 95);t=math.radians(angle);ex=ox+radius*math.sin(t);ey=oy-radius*math.cos(t)
        line(s,ox,oy,ex,ey,RED,4,True);dot(s,ox,oy,TEAL,6);dot(s,ex,ey,RED,7)
        line(s,ox,ey,ex,ey,BLUE,2,False,True);line(s,ox,oy,ox,ey,TEAL,3)
        arc=[(ox+38*math.sin(j*t/50),oy-38*math.cos(j*t/50)) for j in range(51)];poly(s,arc,GOLD,3)
        text(s,f'{angle}°',1120 if n==67 else 1080,455 if n==67 else oy-47,85,40,25,RED,True)
        text(s,'A',ox-30,oy+9,40,35,23,TEAL,True);text(s,'B',ex+10,ey-10,40,35,23,RED,True)
        note='南西へ：ΔX − / ΔY −' if n==68 else ('北東へ：ΔX ＋ / ΔY ＋' if n==67 else '南東へ：ΔX − / ΔY ＋')
        text(s,note,871,574,330,32,21,RED,True,'center')

# 元の本文を作り直したページは、元の説明を発表者ノートに追記。
raw=(SRC/'slides.md').read_text();pages=re.sub(r'^---.*?\n---\s*','',raw,count=1,flags=re.S).split('\n---\n')
for n in redrawn:
    s=p.slides[n-1];original=re.sub(r'<!--.*?-->','',pages[n-1],flags=re.S).strip()
    tf=s.notes_slide.notes_text_frame
    tf.text=tf.text+'\n\n【視覚版 2026-10-20】\n最初に問いを読み、図から予想させてから説明する。\n【改訂前の本文（内容の照合用）】\n'+original

# 日本語本文は実在する日本語フォントへ統一。数式ツリー内は触らない。
for s in p.slides:
    for el in s._element.iter():
        if el.tag in [f'{{{NS["a"]}}}latin',f'{{{NS["a"]}}}ea',f'{{{NS["a"]}}}cs'] and el.get('typeface')=='Meiryo':
            if not any(a.tag==f'{{{NS["m"]}}}oMath' for a in el.iterancestors()):el.set('typeface',FONT)
dest=OUT/'slides_視覚リニューアル_20261020.pptx'
p.save(dest)
# 全Office数式の内容と順序が同じことを照合（名前空間宣言の差は除く）。
def math_signature(file):
    result={}
    with ZipFile(file) as z:
        for n in range(1,82):
            r=E.fromstring(z.read(f'ppt/slides/slide{n}.xml'))
            def canon(el):return (el.tag,tuple(sorted(el.attrib.items())),el.text,tuple(canon(c) for c in el))
            result[n]=[canon(el) for el in r.findall('.//m:oMath',NS)]
    return result
before=math_signature(SRC/'slides.pptx');after=math_signature(dest)
assert before==after,'Office math changed'
assert hashes=={f:hashlib.sha256((SRC/f).read_bytes()).hexdigest() for f in hashes},'Original changed'
report={'slides':len(p.slides),'redrawn_slides':redrawn,'math_diagram_slides':side_diagrams,'native_equations_preserved':sum(map(len,before.values())),'compatibility_fallbacks':fallback_count,'original_sha256':hashes,'source_unchanged':True,'native_math_unchanged':True}
(REVIEW/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2))
