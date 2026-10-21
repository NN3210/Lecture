#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""第3回・改2 の組み立て：改1 の枚を「構成案」の順に並べ替え、使わない枚を削除し、非表示を解除し、
橋渡しの新規枚を「視覚版」の既存枚を型紙として複製して差し込む。PowerPoint は使わない（python-pptx のみ）。

使い方:
    python -I scripts/kai2_build_pptx.py <改1.pptx> <構成案.md> <出力.pptx> [--new-spec <新規枚.md>] [--renumber]

構成案（Markdown 表）の読み方:
    行の 2 列目が「k1:NN」なら改1 の NN 枚目（1 始まり、改1 の並び順）、「new:ID」なら新規枚。
    | 順 | 元 | 題 | 備考 |
    | 1 | k1:1 | 表紙 | |
    | 2 | new:goal | 今日のゴール | 新規 |

新規枚の仕様（Markdown）:
    ## new:goal
    - template: 25            ← 型紙にする改1 の枚（3 枚カード型＝25、本文型＝15 など）
    - header: 測量学 第3回  /  00  |  今日のゴール
    - title: 今日のゴール
    - subtitle: 1 行の副題（省略可）
    - card1: 見出し | 本文 1 行目 / 2 行目      ← card1〜card3（型紙 25 のとき）
    - bullets: 1 つ目 / 2 つ目 / 3 つ目           ← 箇条書き型（カードを消して 1 つの枠にする）
    - takeaway: 下の帯の 1 文
    - notes: 発表者ノート（省略可）
"""
from __future__ import annotations

import argparse
import copy
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches

sys.stdout.reconfigure(encoding="utf-8")

NS_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
ROW = re.compile(r"^\|\s*\d+\s*\|\s*(k1:(\d+)|new:([A-Za-z0-9_\-]+))\s*\|")

# 型紙 25（3 枚カード型）の部品名
T25 = {
    "header": "TextBox 2", "title": "TextBox 3", "subtitle": "TextBox 4", "page": "TextBox 5",
    "cards": [("TextBox 9", "TextBox 10"), ("TextBox 14", "TextBox 15"), ("TextBox 19", "TextBox 20")],
    "card_shapes": ["Rounded Rectangle 6", "Rounded Rectangle 7", "TextBox 8", "TextBox 9", "TextBox 10",
                    "Rounded Rectangle 11", "Rounded Rectangle 12", "TextBox 13", "TextBox 14", "TextBox 15",
                    "Rounded Rectangle 16", "Rounded Rectangle 17", "TextBox 18", "TextBox 19", "TextBox 20"],
    "takeaway": "TextBox 23",
}


def read_plan(path: Path) -> list[tuple[str, str]]:
    plan = []
    for line in path.read_text(encoding="utf-8").splitlines():
        m = ROW.match(line)
        if m:
            plan.append(("k1", m.group(2)) if m.group(2) else ("new", m.group(3)))
    return plan


def read_spec(path: Path) -> dict[str, dict]:
    spec: dict[str, dict] = {}
    cur = None
    for line in path.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^##\s+new:([A-Za-z0-9_\-]+)", line)
        if m:
            cur = spec.setdefault(m.group(1), {})
            continue
        m = re.match(r"^-\s+([a-z0-9]+):\s*(.*)$", line)
        if m and cur is not None:
            cur[m.group(1)] = m.group(2).strip()
    return spec


def shape_by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(name)


def set_text(shape, lines: list[str]) -> None:
    """最初の段落・最初のランの書式を残して、文を入れ替える。"""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    rpr = copy.deepcopy(p0.runs[0]._r.rPr) if p0.runs and p0.runs[0]._r.rPr is not None else None
    ppr = copy.deepcopy(p0._p.pPr) if p0._p.pPr is not None else None
    # 2 段落目以降を消す
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    for i, text in enumerate(lines):
        p = p0 if i == 0 else tf.add_paragraph()
        if i > 0 and ppr is not None:
            p._p.insert(0, copy.deepcopy(ppr))
        r = p.add_run()
        r.text = text
        if rpr is not None:
            r._r.insert(0, copy.deepcopy(rpr))


def clone_slide(prs, src_slide):
    """型紙の枚を複製して末尾に追加する（図は無い前提。関係のコピーはしない）。"""
    layout = src_slide.slide_layout
    new = prs.slides.add_slide(layout)
    for sh in list(new.shapes):  # レイアウトのプレースホルダーを消す
        sh._element.getparent().remove(sh._element)
    for el in src_slide.shapes._spTree:
        if el.tag.endswith("}sp") or el.tag.endswith("}cxnSp") or el.tag.endswith("}grpSp"):
            new.shapes._spTree.append(copy.deepcopy(el))
    return new


def build_new_slide(prs, spec: dict, sid: str):
    tnum = int(spec.get("template", "25"))
    tmpl = prs.slides[tnum - 1]
    s = clone_slide(prs, tmpl)
    names = T25
    if "header" in spec:
        set_text(shape_by_name(s, names["header"]), [spec["header"]])
    if "title" in spec:
        set_text(shape_by_name(s, names["title"]), [spec["title"]])
    if "subtitle" in spec:
        set_text(shape_by_name(s, names["subtitle"]), [spec["subtitle"]])
    elif tnum == 25:
        set_text(shape_by_name(s, names["subtitle"]), [""])
    if tnum == 25:
        if "bullets" in spec:
            # カードを消して、カード 1 の本文枠を広げて箇条書きにする
            body = shape_by_name(s, "TextBox 10")
            keep = {"TextBox 10"}
            for nm in names["card_shapes"]:
                if nm not in keep:
                    shape_by_name(s, nm)._element.getparent().remove(shape_by_name(s, nm)._element)
            body.left, body.top, body.width, body.height = Inches(0.9), Inches(2.3), Inches(11.6), Inches(3.9)
            set_text(body, ["・" + b.strip() for b in spec["bullets"].split(" / ")])
        else:
            cards = [spec.get(f"card{i}") for i in (1, 2, 3)]
            for (hname, bname), c in zip(names["cards"], cards):
                if not c:
                    continue
                head, _, body = c.partition("|")
                set_text(shape_by_name(s, hname), [head.strip()])
                set_text(shape_by_name(s, bname), [ln.strip() for ln in body.split(" / ") if ln.strip()])
    if "takeaway" in spec:
        set_text(shape_by_name(s, names["takeaway"]), [spec["takeaway"]])
    if "notes" in spec:
        s.notes_slide.notes_text_frame.text = spec["notes"]
    s.name = f"new:{sid}"
    return s


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("plan")
    ap.add_argument("dst")
    ap.add_argument("--new-spec", default="")
    ap.add_argument("--renumber", action="store_true", help="「NN / 81」形式のページ番号を新しい番号に書き換える")
    a = ap.parse_args()

    plan = read_plan(Path(a.plan))
    if not plan:
        sys.exit("構成案に k1:/new: の行が無い")
    spec = read_spec(Path(a.new_spec)) if a.new_spec else {}
    prs = Presentation(a.src)
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    n_src = len(ids)

    used = [int(v) for k, v in plan if k == "k1"]
    bad = [n for n in used if n < 1 or n > n_src]
    if bad:
        sys.exit(f"改1 に無い枚番号: {bad}（改1 は {n_src} 枚）")
    dup = sorted({n for n in used if used.count(n) > 1})
    if dup:
        sys.exit(f"同じ枚が 2 回以上使われている: {dup}")
    missing_spec = [v for k, v in plan if k == "new" and v not in spec]
    if missing_spec:
        sys.exit(f"新規枚の仕様が無い: {missing_spec}")

    # 新規枚を作る（末尾に追加される）
    new_ids: dict[str, object] = {}
    for k, v in plan:
        if k == "new":
            build_new_slide(prs, spec[v], v)
            new_ids[v] = list(sldIdLst)[-1]

    # 並べ替え
    order = [ids[int(v) - 1] if k == "k1" else new_ids[v] for k, v in plan]
    for el in list(sldIdLst):
        sldIdLst.remove(el)
    for el in order:
        sldIdLst.append(el)

    # 使わない枚を削除（関係も落とす）
    keep = set(used)
    removed = 0
    for n, el in enumerate(ids, 1):
        if n not in keep:
            prs.part.drop_rel(el.get(NS_R))
            removed += 1

    # 非表示解除
    for slide in prs.slides:
        if slide._element.get("show") == "0":
            del slide._element.attrib["show"]

    # ページ番号の書き換え
    if a.renumber:
        total = len(plan)
        pat = re.compile(r"^\s*\d{1,3}\s*/\s*\d{2,3}\s*$")
        for pos, slide in enumerate(prs.slides, 1):
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for para in shape.text_frame.paragraphs:
                    txt = "".join(r.text for r in para.runs)
                    if pat.match(txt) and para.runs:
                        para.runs[0].text = f"{pos:02d} / {total}"
                        for r in para.runs[1:]:
                            r.text = ""

    prs.save(a.dst)
    print(f"改1 {n_src} 枚 → 使用 {len(used)} 枚、削除 {removed} 枚、新規 {len(new_ids)} 枚、出力 {len(prs.slides)} 枚: {a.dst}")


if __name__ == "__main__":
    main()
