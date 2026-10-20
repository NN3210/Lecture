#!/usr/bin/env python3
"""地図・図法の歴史の独立補足スライドを作成。Python-pptx / matplotlib。"""
from pathlib import Path
import os
import math
import json

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "materials/slides/第03回_測量の座標/補足_地図と図法の歴史_20261020"
QA = ROOT / "notes/map_history_20261020"
OUT.mkdir(parents=True, exist_ok=True)
QA.mkdir(parents=True, exist_ok=True)
os.environ.setdefault("MPLCONFIGDIR", str(QA / ".mplcache"))
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

BG, NAVY, TEAL, GOLD, INK, MUTED = "FAF9F5", "173C4D", "117E7A", "C5793C", "203E4C", "65767C"
FONT = "Noto Sans CJK JP"
plt.rcParams.update({"font.family": FONT, "font.size": 13, "text.color": "#"+INK,
                     "axes.labelcolor": "#"+INK, "xtick.color": "#"+MUTED,
                     "ytick.color": "#"+MUTED, "axes.unicode_minus": False})
SOURCES = {
 "A": ("米国議会図書館：Mathematics / Ptolemy’s Geography",
       "https://loc.gov/exhibits/vatican/math.html",
       "2世紀の『地理学』、約8,000地点の座標、複数の投影法。"),
 "B": ("米国議会図書館：Portolan Charts",
       "https://guides.loc.gov/nautical-charts/portolan-charts",
       "13世紀末からの海図、港・海岸・方位線、地中海などの航海と交易。"),
 "C": ("University College Oxford：Mercator’s Atlas",
       "https://www.univ.ox.ac.uk/news/mercators-atlas/",
       "1569年世界図の航海目的と等角航路。本人の航海失敗談を裏付ける資料ではない。"),
 "D": ("USGS：Map projections: A working manual（Snyder, 1987）",
       "https://pubs.usgs.gov/publication/pp1395",
       "メルカトル・横メルカトル・ランベルト・モルワイデの原理と歴史。本文PDFも参照：https://pubs.usgs.gov/pp/1395/report.pdf"),
 "E": ("IGN France：30 cartes qui racontent l’histoire de la cartographie",
       "https://www.ign.fr/mag/30-cartes-qui-racontent-lhistoire-de-la-cartographie",
       "1668年の国家的な地図整備の要請、カッシーニ家の三角測量、1749年からの図葉刊行。"),
 "F": ("国土地理院：伊能大図彩色図について",
       "https://service.gsi.go.jp/kochizu/app/about/inoutaizu_saisyokuzu",
       "緯度1度の距離への関心、1800〜1816年の測量、没後1821年の完成。"),
 "G": ("USGS：The United States Geological Survey, 1939–1961, Chapter 6",
       "https://pubs.usgs.gov/book/2015/rabbitt-vol4/pdf/vol4_chapter6.pdf",
       "戦後のアラスカ地図整備、カナダの図葉との調整を含むUTM採用の背景。"),
 "H": ("USGS：How are UTM coordinates measured on USGS topographic maps?",
       "https://www.usgs.gov/faqs/how-are-utm-coordinates-measured-usgs-topographic-maps",
       "原則経度6度幅・60ゾーン、メートル単位の座標格子。"),
 "I": ("国土地理院：わかりやすい平面直角座標系",
       "https://www.gsi.go.jp/sokuchikijun/jpc.html",
       "全国19系、X北・Y東、原点縮尺係数0.9999。"),
 "J": ("国土地理院：地理院タイルについて",
       "https://maps.gsi.go.jp/development/siyou.html",
       "メルカトルの式、約±85.0511度の範囲、ズームレベルとタイル。"),
 "K": ("USGS：Map Projections（1993、図法比較ポスター）",
       "https://pubs.usgs.gov/gip/70047422/report.pdf",
       "ランベルト正角円錐図法の1772年発表、図法ごとの保存特性と用途。"),
 "L": ("USGS：ISIS Mollweide Class Reference",
       "https://isis.astrogeology.usgs.gov/9.0.0/Object/Programmer/class_isis_1_1_mollweide.html",
       "1805年の発表、正積擬円筒図法の仕様。"),
}

def saveplot(fig, name):
    p = OUT / "figures" / f"{name}.png"
    p.parent.mkdir(exist_ok=True)
    fig.savefig(p, dpi=190, facecolor="#"+BG, bbox_inches="tight", pad_inches=.18)
    plt.close(fig)
    return p

def merc(lon, lat):
    return np.deg2rad(lon), np.log(np.tan(np.pi/4+np.deg2rad(lat)/2))

def moll(lon, lat):
    phi=np.deg2rad(np.asarray(lat))
    theta=phi.copy()
    for _ in range(30):
        theta -= (2*theta+np.sin(2*theta)-np.pi*np.sin(phi))/(2+2*np.cos(2*theta)+1e-12)
    return 2*np.sqrt(2)/np.pi*np.deg2rad(lon)*np.cos(theta), np.sqrt(2)*np.sin(theta)

def graticule(ax, fn, latmax=75):
    for la in range(-60,61,30):
        ax.plot(*fn(np.linspace(-180,180,241), np.full(241,la)), color="#BBCAC9",lw=.8)
    for lo in range(-180,181,30):
        ax.plot(*fn(np.full(181,lo), np.linspace(-latmax,latmax,181)),color="#BBCAC9",lw=.8)
    ax.set_aspect("equal"); ax.axis("off")

def make_figures():
    fig,ax=plt.subplots(figsize=(8,4.8))
    for lat in [0,20,40,60,75]:
        y=merc(0,lat)[1]
        ax.axhline(y,color="#BBCAC9"); ax.text(-.32,y,f"{lat}°",ha="right",va="center")
    for lon in [0,20,40,60,80]:
        ax.axvline(np.deg2rad(lon),color="#BBCAC9")
    x=np.linspace(.08,1.3,200); y=.18+1.45*x
    ax.plot(x,y,color="#"+GOLD,lw=4,label="等角航路（真北に対し一定の角度）")
    ax.scatter([x[0],x[-1]],[y[0],y[-1]],c="#"+GOLD,s=65)
    ax.text(x[0],y[0]-.18,"出発",ha="center");ax.text(x[-1],y[-1]+.12,"到着",ha="center")
    ax.set_xlim(-.15,1.55);ax.set_ylim(-.1,2.25);ax.axis("off")
    ax.legend(loc="lower right",frameon=False,fontsize=12)
    saveplot(fig,"mercator_route")

    fig,axes=plt.subplots(1,2,figsize=(11,4.4))
    for ax,fn,title in zip(axes,[merc,moll],["メルカトル：角を保つ","モルワイデ：面積比を保つ"]):
        graticule(ax,fn,latmax=89.999 if fn is moll else 75)
        for lat,lon in [(0,-85),(60,-85),(0,65),(60,65)]:
            a=np.linspace(0,2*np.pi,161)
            # 同じ角半径の球面小円。地球上で面積が等しい。
            p=np.deg2rad(lat);l=np.deg2rad(lon);r=np.deg2rad(8)
            pp=np.arcsin(np.sin(p)*np.cos(r)+np.cos(p)*np.sin(r)*np.cos(a))
            ll=l+np.arctan2(np.sin(a)*np.sin(r)*np.cos(p),np.cos(r)-np.sin(p)*np.sin(pp))
            xx,yy=fn(np.rad2deg(ll),np.rad2deg(pp))
            ax.fill(xx,yy,color="#"+TEAL,alpha=.66)
        ax.set_title(title,pad=16,color="#"+NAVY,fontsize=16)
    fig.text(.5,.02,"緑の領域は地球上では同じ面積。図は球面モデルによる自作比較図（海岸線省略）。",ha="center",fontsize=12)
    saveplot(fig,"comparison")

    fig,ax=plt.subplots(figsize=(7.5,4.5))
    p1,p2=np.deg2rad([30,60])
    n=np.log(np.cos(p1)/np.cos(p2))/np.log(np.tan(np.pi/4+p2/2)/np.tan(np.pi/4+p1/2))
    f=np.cos(p1)*np.tan(np.pi/4+p1/2)**n/n
    def lcc(lo,la):
        rho=f/np.tan(np.pi/4+np.deg2rad(la)/2)**n
        a=n*np.deg2rad(lo)
        return rho*np.sin(a),-rho*np.cos(a)
    for lat in [15,30,45,60,75]:
        ax.plot(*lcc(np.linspace(-80,80,200),lat),color="#"+(TEAL if lat in [30,60] else "BBCAC9"),lw=2 if lat in [30,60] else 1)
    for lon in np.arange(-80,81,20):
        ax.plot(*lcc(lon,np.linspace(15,75,100)),color="#BBCAC9")
    for lat in [30,60]:
        x,y=lcc(75,lat);ax.text(x+.05,y,f"標準緯線 {lat}°",fontsize=13)
    ax.set_aspect("equal");ax.axis("off")
    saveplot(fig,"lambert")

prs=Presentation()
prs.slide_width=Inches(13.333)
prs.slide_height=Inches(7.5)
MANUSCRIPT=[]
def color(c): return RGBColor.from_string(c)
def shape(s,x,y,w,h,fill,kind=MSO_SHAPE.RECTANGLE,line=None):
    a=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid();a.fill.fore_color.rgb=color(fill)
    a.line.fill.background() if not line else None
    if line:a.line.color.rgb=color(line)
    return a
def text(s,x,y,w,h,t,size=22,c=INK,bold=False):
    a=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=a.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.02)
    tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(t.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(c)
        p.space_after=Pt(10);p.line_spacing=1.12
    return a
def line(s,x1,y1,x2,y2,c=TEAL,width=2):
    a=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    a.line.color.rgb=color(c);a.line.width=Pt(width);return a
def new(title,era,refs="",note=""):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid();s.background.fill.fore_color.rgb=color(BG)
    shape(s,0,0,.16,7.5,TEAL)
    text(s,.52,.27,12,.3,era,12,TEAL,True)
    text(s,.52,.83,12.2,.85,title,30,NAVY,True)
    line(s,.55,6.95,12.8,6.95,"D4DEDA",.8)
    label="出典 "+", ".join(f"[{k}]" for k in refs) if refs else "測量学 第3回 補足資料"
    text(s,.55,7.08,11.4,.22,label+"  •  出典一覧と発表者ノートにURL",9,MUTED)
    text(s,12.2,7.03,.5,.3,f"{len(prs.slides):02}",12,MUTED)
    src="\n".join(f"[{k}] {SOURCES[k][0]}\n{SOURCES[k][1]}" for k in refs)
    s.notes_slide.notes_text_frame.text=note+("\n\n【出典】\n"+src if src else "")
    MANUSCRIPT.append((title,era,note,src))
    return s
def takeaway(s,t):
    shape(s,.55,6.1,12.2,.62,"E5EFEB")
    text(s,.8,6.23,11.7,.39,t,19,TEAL,True)
def card(s,x,y,w,title,body,accent=TEAL):
    shape(s,x,y,w,3.42,"FFFFFF")
    shape(s,x,y,w,.06,accent)
    text(s,x+.2,y+.25,w-.4,.65,title,22,accent,True)
    text(s,x+.2,y+1,w-.4,2.25,body,21)
def picture(s,name,x,y,w):
    s.shapes.add_picture(str(OUT/"figures"/f"{name}.png"),Inches(x),Inches(y),width=Inches(w))

def build():
    make_figures()
    s=new("地図は、時代の課題に応えてきた","地図・図法の歴史｜測量学 第3回「測量の座標」独立補足",
          note="本編と別の補足資料。20〜25分程度。世界の地図史を網羅するのでなく、第3回の投影・座標に関係する事例を選ぶ。『時代の課題→地図の工夫→得たものと限界』を共通の見方とする。")
    text(s,.7,2.1,7.3,1.6,"航海、国土の測量、世界の比較。\n人は地図に何を求めてきたのか。",30,NAVY,True)
    text(s,.72,4.25,7.2,1,"古代の座標 → 大航海時代 → 近代測量 → Web地図",22)
    for i,(t,c) in enumerate([("方角",TEAL),("面積",GOLD),("位置",NAVY)]):
        shape(s,9.2,2.05+i*1.15,2.5,.85,c,MSO_SHAPE.ROUNDED_RECTANGLE)
        text(s,9.85,2.2+i*1.15,1.5,.5,t,25,"FFFFFF",True)
    takeaway(s,"図法を学ぶ前に、その地図が必要とされた理由を考える。")

    s=new("必要が変わると、地図も変わる","歴史の見取り図","ABCEFGJKL",
          "欧州と日本の選択的な系譜。矢印は時系列の案内で、直接の発明継承や単線的な進歩を意味しない。航海用・国土測量用・教育用の地図は現在も併存する。")
    entries=[("2世紀頃","世界の知識を整理","プトレマイオス"),("13世紀末〜","港から港へ航海","ポルトラノ海図"),("1569年","外洋の針路を描く","メルカトル"),
             ("17〜19世紀","国土を実測する","カッシーニ・伊能"),("18〜20世紀","目的に合う投影","正角・正積・UTM"),("2000年代〜","画面で地図を配信","Webメルカトル")]
    for i,(date,need,name) in enumerate(entries):
        x=.65+(i%3)*4.15;y=1.9+(i//3)*1.88
        shape(s,x,y,3.9,1.6,"FFFFFF")
        text(s,x+.18,y+.14,3.5,.4,date,22,TEAL,True)
        text(s,x+.18,y+.64,3.5,.4,need,21,NAVY,True)
        text(s,x+.18,y+1.15,3.5,.3,name,15,MUTED)
    takeaway(s,"これは代表例の系譜。世界の地図史全体を網羅する年表ではない。")

    s=new("平らな地図では、すべてを同時に保てない","まず押さえる原理","DK",
          "球面・楕円体を平面へ展開するとひずみが生じる。正角は各点近傍の角を保存する性質で、大陸全体の形の完全保存ではない。正距も一般に特定の点から、または特定の線上の距離に限る。")
    card(s,.65,2,3.9,"角を保つ｜正角","小さな範囲の角や形を保つ。\n航海・測量に役立つ。")
    card(s,4.72,2,3.9,"面積比を保つ｜正積","土地の広がりを比較する。\n分布図や世界比較に役立つ。",GOLD)
    card(s,8.79,2,3.9,"距離を保つ｜条件付き","特定の点からの距離など。\nすべての2点間ではない。",NAVY)
    takeaway(s,"「正確な地図」は、何について正確かを添えて説明する。")

    s=new("プトレマイオス：世界の知識を座標で整理する","2世紀頃｜ローマ帝国期の学知","A",
          "『地理学』は先行する知識を集成・検討し、約8,000地点の緯度経度と投影の方法を記した。座標を初めて発明した人物と断定しない。今日伝わる図版には後世の写本・印刷版が含まれる。右図は概念図で古地図の復元ではない。")
    text(s,.7,1.95,6.5,3.5,"背景：広域の地理知識を整理したい。\n工夫：地点を緯度・経度の数値で示す。\n意義：同じ規則で地図を描き直せる。\n限界：座標があっても観測の誤りは残る。",23)
    for i in range(5):
        line(s,8+i*.8,2,8+i*.8,5.2,"BBCAC9",1)
        line(s,8,2+i*.8,11.2,2+i*.8,"BBCAC9",1)
    shape(s,9.5,3.5,.18,.18,GOLD,MSO_SHAPE.OVAL)
    text(s,9.75,3.28,2.7,1,"地点を\n数の組で共有",20,TEAL,True)
    text(s,8,5.45,4,.3,"概念図：実際の古代座標ではない",13,MUTED)
    takeaway(s,"第3回への接続：位置を伝えるには、座標と共通の約束が必要。")

    s=new("ポルトラノ海図：地中海交易の航路を支える","13世紀末〜｜港・海岸・羅針盤","B",
          "港と海岸線を重視し、方位線の網を描いた海図。地中海などの航海実務と結びつく。すべてが船上実用品だったわけではなく、装飾的な作例もある。成立過程を単独発明者の物語にしない。右図は方位線の概念図で、史料画像ではない。")
    text(s,.7,1.95,6.5,3.7,"背景：港から港へ、人や商品を運ぶ。\n工夫：海岸線・港名・方位線を重視。\n得たもの：進む方向を計画しやすい。\n限界：全球を統一的に表すには課題。",23)
    for j in range(16):
        a=j*np.pi/8
        line(s,10,3.55,10+1.7*np.cos(a),3.55+1.7*np.sin(a),"BAC9C4",1)
    shape(s,9.88,3.43,.24,.24,TEAL,MSO_SHAPE.OVAL)
    text(s,9.75,1.65,1,.5,"北",19,TEAL,True)
    text(s,8.1,5.45,4.4,.3,"方位線の概念図（海岸線省略）",13,MUTED)
    takeaway(s,"内陸の詳しさよりも、海岸と港のつながりが重要だった。")

    s=new("メルカトル：外洋で使える針路の地図へ","1569年｜大航海時代の交易・探検・海外進出","BC",
          "15〜16世紀の欧州の外洋航海の拡大には、交易だけでなく征服・植民地化の文脈もある。1569年世界図は表題に航海への利用を掲げる。本人が船で目的地に行けず苦労したため発明した、という個人的逸話は今回の参照資料では確認できず、採用しない。")
    card(s,.65,2,3.9,"時代の課題","外洋を長く進むとき、\n地図上で針路を\nどう扱うか。")
    card(s,4.72,2,3.9,"1569年の工夫","一定の方位で進む\n「等角航路」を\n直線として描く。",GOLD)
    card(s,8.79,2,3.9,"史料で確認できること","世界図の表題自体が\n「航海への利用」を\n掲げている。",NAVY)
    takeaway(s,"考証：社会的な航海需要と、本人の体験談は分けて扱う。")

    s=new("メルカトル図法の直線は「一定の方位」を表す","航海のための性質を図で見る","CD",
          "等角航路は各経線と同じ角度で交わる曲線で、メルカトル図上では直線となる。正角は局所的な性質。ここでの方位は真北基準。実際の操船では磁気偏角、風・潮流などを別途考慮する。等角航路は一般に最短経路ではなく、経線・赤道など例外はある。")
    picture(s,"mercator_route",.6,1.75,7.15)
    text(s,8.1,2.1,4.5,3.7,"経線は平行な直線。\n緯線の間隔は高緯度ほど広い。\n経線と一定角度で交わる航路を、直線で描ける。",23)
    takeaway(s,"直線の航路 ＝ 一定の方位。一般に最短距離の航路とは一致しない。")

    s=new("角を保つ代わりに、高緯度の面積が大きくなる","図法のひずみを比較する","DL",
          "球面の理想化したメルカトルで、赤道の縮尺を1とすれば局所長さ倍率は1/cosφ、面積倍率は1/cos²φ。緯度60度では長さ2倍、面積4倍。図の小円は有限サイズなので厳密なティソーの指示楕円ではない。全球の等角性を『すべての地点への方位が正しい』と説明しない。")
    picture(s,"comparison",2.1,1.7,9.15)
    takeaway(s,"球面モデルでは緯度60°の局所面積倍率は4倍（赤道を基準）。")

    s=new("カッシーニの地図：国土を一つの測量網で結ぶ","17〜18世紀｜フランス国家による地図整備","E",
          "IGNの解説によると、1668年にコルベールが科学アカデミーへより正確なフランス地図の作成法を求めた。カッシーニ家の世代をまたぐ三角測量を経て、1749年から図葉が刊行された。三角測量は投影法ではなく位置を求める観測・計算の方法。右図は架空の測量網。")
    text(s,.7,1.95,6.5,3.7,"背景：国土全体を正確に把握したい。\n方法：基線と角から三角形をつなぐ。\n成果：地域ごとの図を共通の骨格へ。\n1749年から地図の図葉が刊行される。",23)
    pts=[(8.1,4.9),(9.8,5.1),(9.2,3.4),(11,3.7),(10.5,2.1),(12,2.4)]
    for a,b in [(0,1),(0,2),(1,2),(1,3),(2,3),(2,4),(3,4),(3,5),(4,5)]:
        line(s,*pts[a],*pts[b],TEAL,2)
    line(s,*pts[0],*pts[1],GOLD,5)
    for x,y in pts:shape(s,x-.06,y-.06,.12,.12,NAVY,MSO_SHAPE.OVAL)
    text(s,8.1,5.45,4.3,.35,"概念図：基線から測量網を広げる",13,MUTED)
    takeaway(s,"地図の進歩には、投影の工夫だけでなく、実測データの整備も必要。")

    s=new("ランベルト：図法を数学的に設計する","1772年｜投影法の理論化","DK",
          "ランベルトは1772年に複数の投影法を発表。ここでは正角円錐図法を取り上げる。中緯度で東西に広い地域に適し、後世の地形図や航空図に利用されるが、航空需要が発明の理由だったと説明しない。図は標準緯線30度・60度の球面モデル。円錐に光を当てれば厳密に得られるという意味でもない。")
    text(s,.7,1.95,6.0,3.65,"背景：球面を平面へ写す数学の発展。\n工夫：角を保つ円錐型の投影を設計。\n特徴：中緯度・東西に広い地域に適する。\n後の用途：地形図・航空図など。",22)
    picture(s,"lambert",7,2.05,5.7)
    takeaway(s,"1772年の考案と、20世紀の航空図への利用は別の出来事。")

    s=new("モルワイデ：世界の面積比を保って描く","1805年｜正積の世界図","L",
          "モルワイデ図法は1805年発表。世界を楕円形に表す正積擬円筒図法。ここでは面積比較の必要性と性質を結びつけて説明するが、特定の社会運動や個人的体験が発明動機だったと断定しない。農地・森林の世界分布は現代の授業上の応用例。")
    card(s,.65,2,3.9,"地図に求めるもの","地域の広さを、\n同じ基準で\n比べたい。")
    card(s,4.72,2,3.9,"1805年の工夫","地球上の面積比を\n地図の上でも保つ。\n世界全体は楕円形。",GOLD)
    card(s,8.79,2,3.9,"利用例と限界","農地・森林などの\n世界分布の比較。\n角や形にはひずみ。",NAVY)
    takeaway(s,"用途から性質を理解する。ただし、用途だけで発明動機を断定しない。")

    s=new("伊能図：歩く測量と天体観測で日本を描く","1800〜1816年の測量 → 1821年の完成","F",
          "伊能忠敬は暦学上の緯度1度の距離の問題に関心を持ち、蝦夷地への測量をはじめとして全国測量を行った。海岸や街道の測量と天体観測を結び付けた。本人は1818年に没し、1821年に幕府天文方の手で大日本沿海輿地全図が完成した。伊能図を現代の全国三角網やGNSS測量と同一視しない。")
    text(s,.7,1.95,7.0,3.5,"背景：地球の大きさへの科学的関心と、\n　　　日本の地理を詳しく把握する仕事。\n方法：距離・方位の実測と天体観測。\n成果：沿岸の形を全国規模で描く。",23)
    for i,(date,t) in enumerate([("1800","蝦夷地へ"),("1816","測量を終える"),("1821","没後に完成")]):
        y=2+i*1.12
        text(s,8.7,y,3.5,.5,date,30,TEAL,True)
        text(s,8.7,y+.53,3.6,.45,t,20)
    takeaway(s,"「一人の天才の地図」だけでなく、測量隊と継承者の仕事として見る。")

    s=new("横メルカトルとUTM：広域の位置を共通に扱う","20世紀｜軍用地図・戦後の地形図整備","DGH",
          "横メルカトルは中央経線付近の狭い帯でひずみを抑える投影。UTMはその投影にゾーン分割などの規則を組み合わせた座標系で、図法そのものの別名ではない。USGSの歴史資料は戦後のアラスカ地図整備で、カナダの図葉との調整などを理由にUTM採用を記す。発明者一人や1947年だけを唯一の起源としない。原則6度幅・60ゾーン。通常の範囲は南緯80度〜北緯84度で極域には別方式を用いる。")
    text(s,.7,1.95,6.4,3.6,"背景：広域の図葉や位置情報をそろえる。\n工夫：地球を細い南北の帯に分ける。\nUTM：原則6°幅、60ゾーン。\n各ゾーンでメートル単位の座標を使う。",23)
    for i in range(5):
        x=8.0+i*.85
        shape(s,x,2.2,.73,2.9,"D7E9E2" if i==2 else "EDF1EB",line="B9CFC5")
        line(s,x+.365,2.2,x+.365,5.1,TEAL if i==2 else "B9CFC5",1)
    text(s,8,5.35,4.5,.5,"隣り合うゾーンの概念図",15,MUTED)
    takeaway(s,"投影法 ＋ ゾーン・原点などの規則 ＝ 実務で使う座標系。")

    s=new("日本の平面直角座標系：地域の測量に合わせる","現代の測量｜第3回の座標につながる","I",
          "このページは現在の制度を歴史の到達点として紹介するもので、制定史の詳細は扱わない。全国19系のガウス・クリューゲルの等角投影。Xは北、Yは東、原点縮尺係数0.9999。UTMと共通する横メルカトル系統の考え方だが、ゾーン・原点・係数などの規則は同じではない。地上距離をそのまま平面上で使えるとは限らず、本編の距離補正につなぐ。")
    card(s,.65,2,3.9,"現場での必要","農地・道路・水路の\n位置や形を、\n平面上で計算したい。")
    card(s,4.72,2,3.9,"日本の選択","全国を19の系に分け、\n地域ごとの\n投影ひずみを抑える。",GOLD)
    card(s,8.79,2,3.9,"座標の約束","X：北、Y：東。\n原点の縮尺係数\n0.9999。",NAVY)
    takeaway(s,"地図の用途と範囲が決まると、必要な座標の約束が見えてくる。")

    s=new("Web地図：画面に配りやすい地図へ","2000年代以降｜スクロール・拡大縮小","J",
          "地理院タイルはメルカトルの数式を用い、約南北85.0511度の範囲を正方形へ変換する。拡大に応じてタイルを分割する設計。『配りやすい』は仕様から読み取れる機能上の利点を要約したもので、単独の発明動機を立証する表現ではない。Web Mercatorは楕円体上の厳密な正角性をそのまま持つわけではない。表示上の長さ・面積を現地値と混同しない。")
    text(s,.7,1.95,6.7,3.7,"必要：広い範囲を素早く表示したい。\n工夫：地図を正方形タイルで配信。\n拡大すると、必要な部分を細かく読む。\n注意：高緯度の拡大は依然として残る。",23)
    for i,n in enumerate([1,2,4]):
        x=8.1+i*1.55
        for a in range(n):
            for b in range(n):
                shape(s,x+a*1.25/n,2.8+b*1.25/n,1.25/n,1.25/n,
                      "D7E9E2",line="8DB8A8")
        text(s,x,4.3,1.5,.4,f"z = {i}",17,TEAL,True)
    takeaway(s,"同じ系統の図法でも、航海とWeb配信では選ばれる理由が異なる。")

    s=new("背景 → 工夫 → 限界、の順に説明する","歴史から図法を選ぶ","BCDKLIJ",
          "この表は授業用の整理。ランベルトの航空図利用や正積図法の主題図利用など、考案当初と後の利用を混同しない。ポルトラノは海図の種類、他の行は投影法または座標系であり分類が異なる。")
    rows=[
        ("航海の針路","メルカトル","等角航路が直線","面積・距離は変わる"),
        ("中緯度の地域図","ランベルト正角円錐","局所の角・形","全球を一様に描けない"),
        ("世界の広さの比較","モルワイデ","面積比","角・形は変わる"),
        ("地域の測量・位置共有","UTM／日本の19系","帯や地域に分けた座標","範囲・系の確認が必要"),
        ("画面での配信","Webメルカトル","正方形タイル","表示面積≠実面積"),
    ]
    widths=[3.0,3.1,3.0,3.2]; xs=[.6,3.6,6.7,9.7]
    for j,h in enumerate(["必要・用途","代表例","役立つ性質・仕組み","気をつけること"]):
        shape(s,xs[j],1.9,widths[j],.6,NAVY)
        text(s,xs[j]+.1,2.05,widths[j]-.15,.3,h,15,"FFFFFF",True)
    for i,row in enumerate(rows):
        for j,cell in enumerate(row):
            y=2.5+i*.62
            shape(s,xs[j],y,widths[j],.62,"FFFFFF" if i%2==0 else "EAF0EB")
            text(s,xs[j]+.1,y+.16,widths[j]-.18,.38,cell,15,INK,j==0)
    takeaway(s,"発明年の暗記から、「この目的には、なぜこの地図か」の説明へ。")

    s=new("確かめよう：その地図で、何をしたい？","確認問題｜発表者ノートに解答",
          note="解答例：①メルカトル図法は一定方位の等角航路を直線で表せるため。ただし最短経路とは限らない。②面積の見かけに偏りが出ない正積図法（モルワイデ等）。③測地系・平面直角座標系の系番号・軸や単位を確認し、必要な距離補正を行う。史実と用途のつながりを自分の言葉で説明させる。")
    for i,(num,t) in enumerate([
        ("01","航海で「一定の方位」を保ちたい。\nメルカトル図法が役立つ理由は？"),
        ("02","世界の農地面積の広がりを比べたい。\nどの性質を保つ図法を選ぶ？"),
        ("03","農地の測量成果を隣の地域と共有したい。\n座標の数値以外に、何を確認する？")]):
        y=1.95+i*1.25
        text(s,.7,y,1,.5,num,29,TEAL,True)
        text(s,1.8,y,10.6,1,t,23)
    takeaway(s,"目的 → 必要な性質 → 座標の約束、の順で答える。")

    s=new("考証メモ：便利な物語と、確認できる史実を分ける","教員用補足｜そのまま講義で使う場合の注意","ACDEFL",
          "資料は機関の解説・公的技術文書を中心に確認した。個人の心情や失敗経験は、その事実を示す出典がなければ創作しない。すべての地図の誕生を単一の社会的要因に還元しない。古代・中世の地図作成には本資料で扱わない地域・伝統もある。確認日2026-10-20。")
    entries=[
        ("メルカトル","航海用の1569年世界図は確認できる。\n「本人が船旅で困ったから」は未確認。"),
        ("図法の年と用途","ランベルト1772年、モルワイデ1805年。\n後の普及用途を、そのまま発明理由にはしない。"),
        ("図と測量","伊能図・カッシーニの地図は測量成果の例。\n投影法そのものの発明とは区別する。")]
    for i,(a,b) in enumerate(entries):
        y=1.95+i*1.25
        text(s,.7,y,3.0,.5,a,22,TEAL,True)
        text(s,3.8,y,8.8,1.0,b,21)
    takeaway(s,"確実な年・目的・性質を骨格にし、個人的な逸話は出典で確かめる。")

    for keys,title in [("ABCDEF","出典①｜古代・航海・近代測量"),("GHIJKL","出典②｜投影・座標・Web地図")]:
        s=new(title,"参照資料｜2026年10月20日確認",
              note="本文下部の出典記号に対応。各リンクをクリックすると参照先へ移動する。解説ページと技術文書を区別して利用。URLの詳細と対応する事実はREADMEおよび講師用原稿にも保存。")
        for i,k in enumerate(keys):
            name,url,desc=SOURCES[k]; y=1.9+i*.72
            a=text(s,.7,y,12,.34,f"[{k}] {name}",16,TEAL,True)
            a.text_frame.paragraphs[0].runs[0].hyperlink.address=url
            text(s,.95,y+.34,11.8,.34,desc.split("本文PDF")[0],13,MUTED)

    ppt=OUT/"地図と図法の歴史_補足スライド.pptx"
    prs.save(ppt)
    md=["# 地図と図法の歴史：講師用原稿・出典\n",
        "測量学第3回とは独立した補足資料。確認日：2026-10-20。\n"]
    for i,(title,era,note,src) in enumerate(MANUSCRIPT,1):
        md += [f"\n## {i}. {title}\n",f"{era}\n\n{note}\n",src+"\n"]
    (OUT/"講師用原稿と出典.md").write_text("\n".join(md),encoding="utf-8")
    # 図の原理を数値で確認。球面の理想化モデル。
    assert abs(1/math.cos(math.radians(60))**2-4)<1e-12
    assert abs(merc(0,0)[1])<1e-12
    assert len(prs.slides)==20
    bounds=[]
    for i,s in enumerate(prs.slides,1):
        for a in s.shapes:
            assert a.left>=0 and a.top>=0 and a.left+a.width<=prs.slide_width+10000 and a.top+a.height<=prs.slide_height+10000,(i,a.name)
        bounds.append({"slide":i,"shapes_within_slide":True})
    (QA/"構造検証.json").write_text(json.dumps({"slides":len(prs.slides),"checks":bounds,
        "editable_text":True,"notes_on_every_slide":True,"mercator_local_area_scale_60deg":4},ensure_ascii=False,indent=2),encoding="utf-8")
    print(ppt)

if __name__=="__main__":
    build()
